"""Read-only validation for the ranked offline XLSX delivery."""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Mapping
from xml.etree import ElementTree

from openpyxl import load_workbook

from .workbook import WorkbookModel, build_workbook_model


@dataclass(frozen=True)
class DeliveryReport:
    sheet_rows: Mapping[str, int]
    candidate_count: int
    shortlist_count: int
    snapshot_digest: str
    research_unresolved: bool
    screening_counts: Mapping[str, Mapping[str, int]]


def _normalized(value: object) -> object:
    return "" if value is None else value


def _check_package(path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"workbook not found: {path}")
    if not zipfile.is_zipfile(path):
        raise ValueError("workbook is not a valid XLSX ZIP archive")
    with zipfile.ZipFile(path) as archive:
        names = {name.casefold() for name in archive.namelist()}
        forbidden = [
            name for name in names
            if "vbaproject" in name
            or "externallinks/" in name
            or "externalconnections" in name
            or name.endswith("/connections.xml")
        ]
        if forbidden:
            raise ValueError(f"workbook contains an external link, connection, or macro part: {forbidden[0]}")
        for name in archive.namelist():
            if not name.endswith(".rels"):
                continue
            root = ElementTree.fromstring(archive.read(name))
            for relationship in root:
                if relationship.attrib.get("TargetMode", "").casefold() == "external":
                    raise ValueError(f"workbook contains an external relationship in {name}")


def _check_workbook_values(path: Path, model: WorkbookModel) -> dict[str, int]:
    _check_package(path)
    normal = load_workbook(path, read_only=False, data_only=False)
    read_only = load_workbook(path, read_only=True, data_only=False)
    data_only = load_workbook(path, read_only=False, data_only=True)
    try:
        expected_names = list(model.sheets)
        if normal.sheetnames != expected_names:
            missing = sorted(set(expected_names) - set(normal.sheetnames))
            extra = sorted(set(normal.sheetnames) - set(expected_names))
            raise ValueError(f"workbook sheet list differs; missing={missing}, extra={extra}")
        if read_only.sheetnames != expected_names or data_only.sheetnames != expected_names:
            raise ValueError("workbook does not reopen with the required sheets in every load mode")
        if getattr(normal, "_external_links", ()):
            raise ValueError("workbook contains external links")

        if "Owned" not in normal.sheetnames:
            raise ValueError("workbook is missing the required Owned sheet")
        for sheet_name in ("Shortlist", "All candidates"):
            headers = [cell.value for cell in normal[sheet_name][1]]
            if "owned" not in headers:
                raise ValueError(f"{sheet_name} is missing the required owned column")
        owned_headers = [cell.value for cell in normal["Owned"][1]]
        owned_rows = list(normal["Owned"].iter_rows(min_row=2, values_only=True))
        owned_columns = {name: index for index, name in enumerate(owned_headers)}
        expected_owned = {
            row[owned_columns["domain"]]: row
            for row in owned_rows
        }
        expected_owned_count = len(model.sheets["Owned"].rows)
        if len(owned_rows) != expected_owned_count:
            raise ValueError(
                f"Owned sheet row count mismatch: expected {expected_owned_count}, got {len(owned_rows)}"
            )
        if len(expected_owned) != len(owned_rows):
            raise ValueError("Owned sheet contains duplicate domains")
        for sheet_name in ("Shortlist", "All candidates"):
            sheet = normal[sheet_name]
            headers = [cell.value for cell in sheet[1]]
            columns = {name: index for index, name in enumerate(headers)}
            required_columns = {"owned"}
            if sheet_name == "All candidates":
                required_columns.add("registry_status")
            missing_status = required_columns - set(columns)
            if missing_status:
                raise ValueError(f"{sheet_name} is missing ownership/status columns: {sorted(missing_status)}")
            for row in sheet.iter_rows(min_row=2, values_only=True):
                domain = row[columns["domain"]]
                is_owned = domain in expected_owned
                if (row[columns["owned"]] == "yes") != is_owned:
                    raise ValueError(f"{sheet_name} owned flag does not match Owned sheet for {domain}")
                if is_owned and sheet_name == "All candidates":
                    owned_row = expected_owned[domain]
                    purchase_date = owned_row[owned_columns["purchase_date"]]
                    wording = f"owned by you (bought {purchase_date})"
                    if row[columns["registry_status"]] != wording:
                        raise ValueError(f"All candidates availability wording mismatch for {domain}")

        row_counts: dict[str, int] = {}
        for name, expected_sheet in model.sheets.items():
            sheet = normal[name]
            if sheet.sheet_state != "visible":
                raise ValueError(f"required worksheet is hidden: {name}")
            if sheet.max_row != len(expected_sheet.rows) + 1:
                raise ValueError(f"{name} row count mismatch: expected {len(expected_sheet.rows)}, got {sheet.max_row - 1}")
            if sheet.max_column != len(expected_sheet.headers):
                raise ValueError(f"{name} column count mismatch")
            row_counts[name] = len(expected_sheet.rows)
            expected_grid = (expected_sheet.headers, *expected_sheet.rows)
            normal_rows = sheet.iter_rows()
            read_only_rows = read_only[name].iter_rows()
            data_only_rows = data_only[name].iter_rows()
            for row_number, expected_row in enumerate(expected_grid, start=1):
                normal_row = next(normal_rows)
                read_only_row = next(read_only_rows)
                data_only_row = next(data_only_rows)
                for column_number, expected in enumerate(expected_row, start=1):
                    cell = normal_row[column_number - 1]
                    coord = cell.coordinate
                    if cell.data_type == "f":
                        raise ValueError(f"formula cell is forbidden: {name}!{coord}")
                    if cell.hyperlink is not None:
                        raise ValueError(f"runnable or automatic hyperlink is forbidden: {name}!{coord}")
                    actual = _normalized(cell.value)
                    if expected == "" and actual in (None, ""):
                        for reopened in (
                            read_only_row[column_number - 1],
                            data_only_row[column_number - 1],
                        ):
                            if _normalized(reopened.value) != actual:
                                raise ValueError(f"{name}!{coord} differs across openpyxl load modes")
                        continue
                    if isinstance(expected, str):
                        if not isinstance(actual, str) or actual != expected:
                            raise ValueError(
                                f"{name}!{coord} text mismatch: expected {expected!r}, got {actual!r}"
                            )
                        if cell.data_type != "s":
                            raise ValueError(f"external text is not stored as a plain string: {name}!{coord}")
                    elif expected is None:
                        if actual not in (None, ""):
                            raise ValueError(f"{name}!{coord} must be blank, got {actual!r}")
                    else:
                        if isinstance(actual, bool) or not isinstance(actual, (int, float)):
                            raise ValueError(f"{name}!{coord} must be numeric, got {actual!r}")
                        if actual != expected:
                            raise ValueError(
                                f"{name}!{coord} numeric mismatch: expected {expected!r}, got {actual!r}"
                            )
                    for reopened in (
                        read_only_row[column_number - 1],
                        data_only_row[column_number - 1],
                    ):
                        if _normalized(reopened.value) != actual:
                            raise ValueError(f"{name}!{coord} differs across openpyxl load modes")
        return row_counts
    finally:
        normal.close()
        read_only.close()
        data_only.close()


def validate_delivery(
    workbook_path: str | Path,
    paths: Mapping[str, str | Path] | None = None,
) -> DeliveryReport:
    """Check every workbook cell against the current sources without writing files."""

    model = build_workbook_model(paths)
    workbook = Path(workbook_path).resolve()
    row_counts = _check_workbook_values(workbook, model)
    return DeliveryReport(
        sheet_rows=row_counts,
        candidate_count=int(model.metrics["candidates"]),
        shortlist_count=int(model.metrics["shortlist"]),
        snapshot_digest=model.snapshot_digest,
        research_unresolved=model.research_unresolved,
        screening_counts=model.screening_counts,
    )
