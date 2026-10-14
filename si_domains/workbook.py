"""Build a deterministic, offline XLSX delivery from the checked CSV snapshot."""

from __future__ import annotations

import csv
import hashlib
import os
import tempfile
import textwrap
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from statistics import mean, median, pstdev
from typing import Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import rules
from .config import DATA_DIR, OUTPUT_DIR, PROJECT_ROOT
from .evidence import FIELDNAMES as RDAP_FIELDS
from .evidence import WHOIS_FIELDNAMES
from .research import OPPORTUNITY_FIELDS, _decimal, _unique_by_domain
from .scoring import (
    COMMERCIAL_ROOTS,
    SCORE_FIELDS,
    SCORE_WEIGHTS,
    TRADEMARK_CLEAR_SCORE,
    TRADEMARK_CONFLICT_SCORE,
    TRADEMARK_UNVERIFIED_SCORE,
    _negative_connotations,
    trademark_score_for_status,
    weighted_total,
)
from .shortlist import EXCLUDED_FIELDS, SHORTLIST_FIELDS, SHORTLIST_LIMIT, CATEGORY_LIMIT, select_eligible
from .trademarks import (
    REQUIRED_SOURCES,
    TRADEMARK_EVIDENCE_FIELDS,
    TRADEMARK_SUMMARY_FIELDS,
    aggregate_trademark_rows,
    query_coverage_by_source,
    read_candidate_context,
    read_trademark_evidence,
    read_trademark_summary,
)
from .valuation import (
    VALUATION_EVIDENCE_FIELDS,
    VALUATION_FIELDS,
    eligible_comps,
    read_valuation_evidence,
)


DEFAULT_INPUT_PATHS: dict[str, Path] = {
    "candidates": DATA_DIR / "candidates.csv",
    "reserved_names": DATA_DIR / "reserved_names.csv",
    "restricted_names": DATA_DIR / "restricted_names.csv",
    "rdap_evidence": OUTPUT_DIR / "rdap_evidence.csv",
    "whois_evidence": OUTPUT_DIR / "whois_evidence.csv",
    "scores": OUTPUT_DIR / "scores.csv",
    "shortlist": OUTPUT_DIR / "shortlist.csv",
    "excluded": OUTPUT_DIR / "excluded.csv",
    "opportunities": OUTPUT_DIR / "opportunities.csv",
    "valuations": OUTPUT_DIR / "valuations.csv",
    "valuation_evidence": OUTPUT_DIR / "valuation_evidence.csv",
    "trademark_summary": OUTPUT_DIR / "trademark_summary.csv",
    "trademark_evidence": OUTPUT_DIR / "trademark_evidence.csv",
    "sample": OUTPUT_DIR / "sample_404_for_review.csv",
}

_INPUT_ORDER = tuple(DEFAULT_INPUT_PATHS)
_RDAP_CLASSES = frozenset(
    {"registered", "potentially_unregistered", "unverified", "invalid_syntax", "reserved", "restricted", "premium"}
)
_WHOIS_CLASSES = frozenset({"not_registered", "registered", "unverified"})
_SPECIAL_REASONS = {
    "invalid_syntax": "invalid",
    "reserved": "reserved",
    "restricted": "restricted",
    "premium": "premium",
}
_SCORE_COMPONENTS = tuple(SCORE_WEIGHTS)
_MAX_EXCEL_ROW_HEIGHT = 409
_MAX_ESTIMATED_WRAP_LINES = _MAX_EXCEL_ROW_HEIGHT // 15
_SHORTLIST_EXPORT_PRIORITY = ("rank", "domain", "total", "trademark", "trademark_status")
_SHORTLIST_EXPORT_FIELDS = (
    *_SHORTLIST_EXPORT_PRIORITY,
    *(field for field in OPPORTUNITY_FIELDS if field not in _SHORTLIST_EXPORT_PRIORITY),
    "rdap_http_status",
    "rdap_classification",
    "whois_classification",
)


@dataclass(frozen=True)
class CSVTable:
    path: Path
    headers: tuple[str, ...]
    rows: tuple[dict[str, str], ...]


@dataclass(frozen=True)
class SheetData:
    headers: tuple[str, ...]
    rows: tuple[tuple[object, ...], ...]
    freeze_panes: str | None = "A2"
    filters: bool = True
    wrap_columns: frozenset[str] = frozenset()
    integer_columns: frozenset[str] = frozenset()
    decimal_columns: frozenset[str] = frozenset()


@dataclass(frozen=True)
class WorkbookModel:
    sheets: Mapping[str, SheetData]
    input_paths: Mapping[str, Path]
    input_hashes: Mapping[str, str]
    snapshot_digest: str
    metrics: Mapping[str, object]
    research_unresolved: bool
    screening_counts: Mapping[str, Mapping[str, int]]


def _resolved_inputs(paths: Mapping[str, str | Path] | None) -> dict[str, Path]:
    result = dict(DEFAULT_INPUT_PATHS)
    if paths:
        unknown = set(paths) - set(result)
        if unknown:
            raise ValueError(f"unknown delivery input name: {sorted(unknown)[0]}")
        result.update({name: Path(path) for name, path in paths.items()})
    return {name: path.resolve() for name, path in result.items()}


def _read_csv(path: Path, required: set[str], description: str) -> CSVTable:
    if not path.is_file():
        raise FileNotFoundError(f"{description} CSV not found: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames
        if not headers or any(not header for header in headers):
            raise ValueError(f"{description} CSV has no valid header")
        if len(headers) != len(set(headers)):
            raise ValueError(f"{description} CSV has duplicate column names")
        missing = required - set(headers)
        if missing:
            raise ValueError(
                f"{description} CSV is missing required columns: {', '.join(sorted(missing))}"
            )
        rows: list[dict[str, str]] = []
        for line_number, row in enumerate(reader, start=2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"{description} CSV has a malformed row at line {line_number}")
            rows.append({key: value for key, value in row.items() if key is not None})
    return CSVTable(path, tuple(headers), tuple(rows))


def _index_unique(
    rows: Sequence[dict[str, str]], key: str, description: str, *, normalize_domain: bool = False
) -> dict[str, dict[str, str]]:
    indexed: dict[str, dict[str, str]] = {}
    for row in rows:
        value = row.get(key, "")
        normalized = value.strip().casefold() if normalize_domain else value
        if not normalized or normalized in indexed:
            raise ValueError(f"{description} has a blank or duplicate {key}: {value!r}")
        indexed[normalized] = row
    return indexed


def _domain(value: str, description: str) -> str:
    domain = value.strip().casefold()
    if not domain.endswith(".si") or not domain[:-3]:
        raise ValueError(f"invalid .si domain in {description}: {value!r}")
    return domain


def _utc_timestamp(value: str, description: str) -> datetime:
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"invalid UTC timestamp in {description}: {value!r}") from exc
    if parsed.tzinfo is None or parsed.utcoffset() != timezone.utc.utcoffset(parsed):
        raise ValueError(f"timestamp must be UTC in {description}: {value!r}")
    return parsed


def _timestamp_range(rows: Sequence[dict[str, str]], column: str) -> str:
    values = [row.get(column, "") for row in rows if row.get(column, "")]
    if not values:
        return "not recorded"
    for value in values:
        _utc_timestamp(value, column)
    return f"{min(values)} to {max(values)}"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _display_path(path: Path) -> str:
    try:
        return path.resolve().relative_to(PROJECT_ROOT.resolve()).as_posix()
    except ValueError:
        return str(path)


def _rule_class(label: str, reserved: dict[str, dict[str, str]], restricted: dict[str, dict[str, str]]) -> str:
    return rules.classify(label, reserved_rules=reserved, restricted_rules=restricted)


def _rule_explanation(
    label: str,
    classification: str,
    reserved: dict[str, dict[str, str]],
    restricted: dict[str, dict[str, str]],
) -> tuple[str, str, str]:
    actual, rule_id, source_url = rules.explain(
        label, reserved_rules=reserved, restricted_rules=restricted
    )
    if actual != classification:
        raise ValueError(f"registry rule explanation mismatch for {label}")
    canonical = label.strip().casefold()
    entry = reserved.get(canonical) or restricted.get(canonical)
    if entry:
        explanation = entry.get("note", "").strip()
        if not explanation:
            raise ValueError(f"registry rule inventory lacks an explanation for {canonical}")
        return rule_id, source_url, explanation
    if classification == "invalid_syntax":
        return rule_id, source_url, "Label does not satisfy the known .si syntax rules."
    if classification == "premium":
        return rule_id, source_url, "A source-backed premium rule applies."
    raise ValueError(f"no explanation is available for special class {classification!r}")


def _read_inputs(paths: Mapping[str, Path]) -> dict[str, CSVTable]:
    required: dict[str, set[str]] = {
        "candidates": {"domain_label", "category", "seed"},
        "reserved_names": {"label", "class", "source_url", "note"},
        "restricted_names": {"label", "class", "source_url", "note"},
        "rdap_evidence": set(RDAP_FIELDS),
        "whois_evidence": set(WHOIS_FIELDNAMES),
        "scores": set(SCORE_FIELDS),
        "shortlist": set(SHORTLIST_FIELDS),
        "excluded": set(EXCLUDED_FIELDS),
        "opportunities": set(OPPORTUNITY_FIELDS),
        "valuations": set(VALUATION_FIELDS),
        "valuation_evidence": set(VALUATION_EVIDENCE_FIELDS),
        "trademark_summary": set(TRADEMARK_SUMMARY_FIELDS),
        "trademark_evidence": set(TRADEMARK_EVIDENCE_FIELDS) - {"proof_kind", "proof_path"},
        "sample": {"domain", "category", "rdap_timestamp", "whois_result", "whois_timestamp"},
    }
    return {
        name: _read_csv(paths[name], required[name], name.replace("_", " "))
        for name in _INPUT_ORDER
    }


def _validate_registry_rows(
    candidates: Mapping[str, dict[str, str]],
    rdap_rows: Sequence[dict[str, str]],
    whois_rows: Sequence[dict[str, str]],
) -> tuple[dict[str, dict[str, str]], dict[str, dict[str, str]]]:
    rdap_latest: dict[str, dict[str, str]] = {}
    for row in rdap_rows:
        domain = _domain(row["domain"], "RDAP evidence")
        if domain not in candidates:
            raise ValueError(f"RDAP evidence contains a domain outside candidates: {domain}")
        if row["category"] != candidates[domain]["category"] and not _is_reclassified_evidence_category(
            row["category"], candidates[domain]
        ):
            raise ValueError(f"RDAP category does not match candidate metadata for {domain}")
        if row["classification"] not in _RDAP_CLASSES:
            raise ValueError(f"unknown RDAP classification for {domain}: {row['classification']!r}")
        if row["http_status"]:
            if not row["http_status"].isdigit() or len(row["http_status"]) != 3:
                raise ValueError(f"invalid RDAP HTTP status for {domain}")
            if row["classification"] == "registered" and row["http_status"] != "200":
                raise ValueError(f"registered RDAP result must have HTTP 200 for {domain}")
            if row["classification"] == "potentially_unregistered" and row["http_status"] != "404":
                raise ValueError(f"potentially unregistered RDAP result must have HTTP 404 for {domain}")
        elif row["classification"] in {"registered", "potentially_unregistered"}:
            raise ValueError(f"classified RDAP result is missing HTTP status for {domain}")
        if not row["attempts"].isdigit() or int(row["attempts"]) < 1:
            raise ValueError(f"invalid RDAP attempt count for {domain}")
        _utc_timestamp(row["utc_timestamp"], f"RDAP evidence for {domain}")
        rdap_latest[domain] = row

    whois_latest: dict[str, dict[str, str]] = {}
    for row in whois_rows:
        domain = _domain(row["domain"], "WHOIS evidence")
        if domain not in candidates:
            raise ValueError(f"WHOIS evidence contains a domain outside candidates: {domain}")
        if row["category"] != candidates[domain]["category"] and not _is_reclassified_evidence_category(
            row["category"], candidates[domain]
        ):
            raise ValueError(f"WHOIS category does not match candidate metadata for {domain}")
        if row["classification"] not in _WHOIS_CLASSES:
            raise ValueError(f"unknown WHOIS classification for {domain}: {row['classification']!r}")
        _utc_timestamp(row["utc_timestamp"], f"WHOIS evidence for {domain}")
        whois_latest[domain] = row
    return rdap_latest, whois_latest


def _is_reclassified_evidence_category(
    evidence_category: str, candidate: Mapping[str, str]
) -> bool:
    """Allow historical registry tags to retain the old dictionary category."""

    tags = {tag.strip().casefold() for tag in candidate.get("tags", "").split(";")}
    return (
        evidence_category == "dictionary"
        and candidate.get("category") == "invented"
        and candidate.get("method")
        in {"given-name-reclassified", "technical-term-reclassified"}
        and bool(tags & {"given_name", "technical_term"})
    )


def _score_index(
    rows: Sequence[dict[str, str]],
    candidates: Mapping[str, dict[str, str]],
    rdap_latest: Mapping[str, dict[str, str]],
    summary: Mapping[str, dict[str, str]],
) -> dict[str, dict[str, str]]:
    scores = _index_unique(rows, "domain", "score", normalize_domain=True)
    expected_domains = {
        domain
        for domain, row in rdap_latest.items()
        if row["classification"] == "potentially_unregistered"
    }
    if set(scores) != expected_domains:
        raise ValueError("score domain set must exactly match latest RDAP-404 candidates")
    for domain, row in scores.items():
        if domain not in candidates:
            raise ValueError(f"score contains a domain outside candidates: {domain}")
        if row["category"] != candidates[domain]["category"]:
            raise ValueError(f"score category does not match candidate for {domain}")
        components: dict[str, int] = {}
        for name in _SCORE_COMPONENTS:
            try:
                components[name] = int(row[name])
            except ValueError as exc:
                raise ValueError(f"invalid {name} score for {domain}") from exc
            if not 0 <= components[name] <= 100:
                raise ValueError(f"{name} score is outside 0..100 for {domain}")
        expected_trademark = (
            int(summary[domain]["trademark_score"])
            if domain in summary
            else trademark_score_for_status("not yet screened")
        )
        if components["trademark"] != expected_trademark:
            raise ValueError(f"trademark score does not match research status for {domain}")
        if _decimal(row["total"], f"score total for {domain}") != weighted_total(components):
            raise ValueError(f"weighted score total mismatch for {domain}")
    return scores


def _reason_for(
    label: str,
    rule_class: str,
    rdap_row: Mapping[str, str] | None,
    whois_row: Mapping[str, str] | None,
) -> str | None:
    if rule_class != "ok":
        return _SPECIAL_REASONS.get(rule_class, "invalid")
    rdap_class = (rdap_row or {}).get("classification", "")
    if rdap_class == "registered":
        return "rdap_registered"
    if rdap_class in _SPECIAL_REASONS:
        return _SPECIAL_REASONS[rdap_class]
    if rdap_class != "potentially_unregistered":
        return "rdap_unchecked" if not rdap_class else "rdap_unverified"
    whois_class = (whois_row or {}).get("classification", "")
    if whois_class == "registered":
        return "whois_registered_conflict"
    if whois_class == "unverified":
        return "whois_unverified"
    if whois_class != "not_registered":
        return "whois_unchecked" if not whois_class else "whois_unverified"
    return None


def _validate_shortlist_and_excluded(
    candidate_rows: Sequence[dict[str, str]],
    candidates: Mapping[str, dict[str, str]],
    rdap_latest: Mapping[str, dict[str, str]],
    whois_latest: Mapping[str, dict[str, str]],
    scores: Mapping[str, dict[str, str]],
    shortlist_rows: Sequence[dict[str, str]],
    excluded_rows: Sequence[dict[str, str]],
    reserved: dict[str, dict[str, str]],
    restricted: dict[str, dict[str, str]],
) -> tuple[dict[str, dict[str, str]], dict[str, dict[str, str]], set[str], dict[str, str]]:
    eligible: list[tuple[str, dict[str, str], dict[str, str], object]] = []
    expected_excluded: list[dict[str, str]] = []
    special_by_domain: dict[str, str] = {}
    for candidate in candidate_rows:
        label = candidate["domain_label"].strip().casefold()
        domain = f"{label}.si"
        rdap_row = rdap_latest.get(domain)
        whois_row = whois_latest.get(domain)
        rule_class = _rule_class(label, reserved, restricted)
        special_by_domain[domain] = rule_class
        reason = _reason_for(label, rule_class, rdap_row, whois_row)
        if reason:
            expected_excluded.append(
                {
                    "domain": domain,
                    "category": candidate["category"],
                    "reason": reason,
                    "rdap_classification": (rdap_row or {}).get("classification", ""),
                    "whois_classification": (whois_row or {}).get("classification", ""),
                    "rdap_timestamp": (rdap_row or {}).get("utc_timestamp", ""),
                    "whois_timestamp": (whois_row or {}).get("utc_timestamp", ""),
                }
            )
            continue
        score = scores.get(domain)
        if score is None:
            raise ValueError(f"eligible domain {domain} is missing from scores")
        eligible.append((domain, candidate, score, None))
    selected, selection_reasons = select_eligible(eligible)
    if len(shortlist_rows) > SHORTLIST_LIMIT:
        raise ValueError(f"shortlist exceeds the {SHORTLIST_LIMIT}-domain cap")
    if any(count > CATEGORY_LIMIT for count in Counter(row["category"] for row in shortlist_rows).values()):
        raise ValueError(f"shortlist exceeds the {CATEGORY_LIMIT}-domain category cap")
    expected_domains = [item[0] for item in selected]
    actual_domains = [_domain(row["domain"], "shortlist") for row in shortlist_rows]
    if actual_domains != expected_domains:
        raise ValueError("shortlist membership/order does not match score ranking, category and family caps")
    if [row["rank"] for row in shortlist_rows] != [str(rank) for rank in range(1, len(selected) + 1)]:
        raise ValueError("shortlist ranks must be contiguous and match row order")

    shortlist = _index_unique(shortlist_rows, "domain", "shortlist", normalize_domain=True)
    for rank, (domain, candidate, score, _payload) in enumerate(selected, start=1):
        row = shortlist[domain]
        rdap_row = rdap_latest[domain]
        whois_row = whois_latest[domain]
        expected = {
            "rank": str(rank),
            "domain": domain,
            "category": candidate["category"],
            "total": score["total"],
            **{name: score[name] for name in _SCORE_COMPONENTS},
            "rdap_timestamp": rdap_row["utc_timestamp"],
            "whois_timestamp": whois_row["utc_timestamp"],
            "rationale": score["rationale"],
        }
        if any(row.get(field, "") != value for field, value in expected.items()):
            raise ValueError(f"shortlist fields do not match current evidence and score for {domain}")

    expected_excluded.sort(key=lambda row: row["domain"])
    actual_excluded = list(excluded_rows)
    if actual_excluded != expected_excluded:
        raise ValueError("excluded rows do not match every current ineligible candidate")
    excluded = _index_unique(actual_excluded, "domain", "excluded", normalize_domain=True)
    return shortlist, excluded, set(expected_domains), selection_reasons


def _validate_research_tables(
    *,
    tables: Mapping[str, CSVTable],
    candidates: Mapping[str, dict[str, str]],
    candidate_rows: Sequence[dict[str, str]],
    rdap_latest: Mapping[str, dict[str, str]],
    whois_latest: Mapping[str, dict[str, str]],
    scores: Mapping[str, dict[str, str]],
    shortlist: Mapping[str, dict[str, str]],
    shortlisted_domains: set[str],
    reserved: dict[str, dict[str, str]],
    restricted: dict[str, dict[str, str]],
) -> tuple[dict[str, dict[str, str]], dict[str, dict[str, str]], bool, dict[str, dict[str, int]]]:
    from .research import _screening_counts

    valuation_evidence_rows, valuation_evidence = read_valuation_evidence(
        tables["valuation_evidence"].path
    )
    trademark_evidence_rows, trademark_evidence = read_trademark_evidence(
        tables["trademark_evidence"].path
    )
    candidate_map, query_manifest = read_candidate_context(tables["candidates"].path)
    if set(candidate_map) != set(candidates):
        raise ValueError("candidate context does not match the complete candidate table")
    summary_map = read_trademark_summary(
        tables["trademark_summary"].path,
        evidence_path=tables["trademark_evidence"].path,
        candidate_domains=set(candidates),
        candidates_path=tables["candidates"].path,
    )
    summary_rows = tables["trademark_summary"].rows
    summary_rows_by_domain = _index_unique(summary_rows, "domain", "trademark summary", normalize_domain=True)
    summary_expected_domains = [
        _domain(row["domain"], "shortlist")
        for row in tables["shortlist"].rows[: min(100, len(tables["shortlist"].rows))]
    ]
    summary_domains = [_domain(row["domain"], "trademark summary") for row in summary_rows]
    if summary_domains[:len(summary_expected_domains)] != summary_expected_domains:
        raise ValueError("trademark summary must begin with the final ranked top 100 in rank order")
    historical_domains = summary_domains[len(summary_expected_domains):]
    per_name_evidence_domains = {
        row["domain"].casefold()
        for row in trademark_evidence_rows
        if row["domain"] != "*"
    }
    if historical_domains != sorted(historical_domains) or any(
        domain not in per_name_evidence_domains for domain in historical_domains
    ):
        raise ValueError("historical trademark summaries must be sorted and backed by per-name evidence")
    recomputed_summaries = aggregate_trademark_rows(
        trademark_evidence_rows,
        summary_domains,
        expected_queries_by_domain=query_manifest,
        proof_root=tables["trademark_evidence"].path.parent,
    )
    for expected in recomputed_summaries:
        actual = summary_rows_by_domain[expected["domain"]]
        if any(actual.get(field, "") != expected.get(field, "") for field in TRADEMARK_SUMMARY_FIELDS):
            raise ValueError(f"trademark summary does not match its evidence for {expected['domain']}")

    _score_index(tables["scores"].rows, candidates, rdap_latest, summary_map)
    valuations = _index_unique(
        tables["valuations"].rows, "domain", "valuation", normalize_domain=True
    )
    opportunities = _index_unique(
        tables["opportunities"].rows, "domain", "opportunities", normalize_domain=True
    )
    if set(valuations) != shortlisted_domains or set(opportunities) != shortlisted_domains:
        raise ValueError("shortlist, valuations, and opportunities domain sets must match exactly")

    eligible_ids_by_domain: dict[str, set[str]] = {}
    for domain in shortlisted_domains:
        eligible_ids_by_domain[domain] = {
            row["evidence_id"]
            for row in eligible_comps(domain, shortlist[domain]["category"], valuation_evidence_rows)
        }

    for rank, shortlist_row in enumerate(tables["shortlist"].rows, start=1):
        domain = _domain(shortlist_row["domain"], "shortlist")
        valuation = valuations[domain]
        opportunity = opportunities[domain]
        if valuation.get("rank") != str(rank):
            raise ValueError(f"valuation rank does not match shortlist for {domain}")
        status = valuation.get("status", "")
        refs = [item for item in valuation.get("evidence_ids", "").split(";") if item]
        if any(ref not in valuation_evidence for ref in refs):
            raise ValueError(f"unknown valuation evidence reference for {domain}")
        if status == "insufficiently supported":
            if any(valuation.get(field, "") for field in ("currency", "low", "high")):
                raise ValueError(f"unsupported valuation must have blank currency and price cells for {domain}")
        elif status == "supported_speculative_range":
            if valuation.get("currency") not in {"GBP", "USD"} or not refs:
                raise ValueError(f"supported valuation lacks currency or evidence for {domain}")
            if not set(refs).issubset(eligible_ids_by_domain[domain]):
                raise ValueError(f"supported valuation cites an ineligible comparable for {domain}")
            comp_rows = [valuation_evidence[ref] for ref in refs]
            if any(item["evidence_type"] != "completed_sale" or item["currency"] != valuation["currency"] for item in comp_rows):
                raise ValueError(f"supported valuation cites a non-sale or mixed currency for {domain}")
            if len({item["sold_domain"] for item in comp_rows}) < 2:
                raise ValueError(f"supported valuation needs two distinct sales for {domain}")
            low = _decimal(valuation["low"], f"valuation low for {domain}")
            high = _decimal(valuation["high"], f"valuation high for {domain}")
            if low <= 0 or high < low:
                raise ValueError(f"invalid numeric valuation range for {domain}")
        else:
            raise ValueError(f"invalid valuation status for {domain}: {status!r}")

        candidate = candidates[domain]
        score = scores[domain]
        summary = summary_map.get(domain)
        expected_opportunity = {
            **{field: shortlist_row.get(field, "") for field in OPPORTUNITY_FIELDS[:13]},
            "valuation_status": valuation["status"],
            "valuation_currency": valuation["currency"],
            "valuation_low": valuation["low"],
            "valuation_high": valuation["high"],
            "valuation_evidence_ids": valuation["evidence_ids"],
            "valuation_rationale": valuation["rationale"],
            "trademark_status": summary["status"] if summary else "not yet screened",
            "trademark_coverage": summary.get("coverage", "") if summary else "",
            "trademark_evidence_ids": summary.get("evidence_ids", "") if summary else "",
        }
        if any(opportunity.get(field, "") != value for field, value in expected_opportunity.items()):
            raise ValueError(f"opportunity research joins or provenance do not match for {domain}")
        if shortlist_row["category"] != candidate["category"] or shortlist_row["total"] != score["total"]:
            raise ValueError(f"shortlist score/candidate join mismatch for {domain}")

    top_rows = sorted(tables["shortlist"].rows, key=lambda row: int(row["rank"]))[:100]
    screening_counts, unresolved = _screening_counts(
        top_rows,
        trademark_evidence_rows,
        query_manifest,
        tables["trademark_evidence"].path.parent,
    )
    return valuations, opportunities, unresolved, screening_counts


def _excel_numeric(value: str, column: str, domain: str) -> int | float | None:
    if value == "":
        return None
    if column == "rank" or column in _SCORE_COMPONENTS:
        try:
            return int(value)
        except ValueError as exc:
            raise ValueError(f"invalid integer cell {column} for {domain}") from exc
    try:
        number = Decimal(value)
    except InvalidOperation as exc:
        raise ValueError(f"invalid numeric cell {column} for {domain}") from exc
    if not number.is_finite():
        raise ValueError(f"non-finite numeric cell {column} for {domain}")
    return float(number)


def _candidate_status(
    domain: str,
    rule_class: str,
    rdap_row: Mapping[str, str] | None,
    whois_row: Mapping[str, str] | None,
    selected: set[str],
) -> str:
    if rule_class != "ok":
        return f"special_{rule_class}"
    rdap_class = (rdap_row or {}).get("classification", "")
    if not rdap_class:
        return "unknown"
    if rdap_class == "registered":
        return "registered"
    if rdap_class == "unverified":
        return "unverified"
    if rdap_class != "potentially_unregistered":
        return f"special_{rdap_class}"
    whois_class = (whois_row or {}).get("classification", "")
    if whois_class == "registered":
        return "whois_registered_conflict"
    if whois_class == "unverified":
        return "whois_unverified"
    if not whois_class:
        return "whois_unchecked"
    if whois_class != "not_registered":
        return "unknown"
    return "shortlisted" if domain in selected else "eligible_outside_shortlist_cap"


def _metric_counts(
    candidate_rows: Sequence[dict[str, str]],
    rdap_latest: Mapping[str, dict[str, str]],
    whois_latest: Mapping[str, dict[str, str]],
    scores: Mapping[str, dict[str, str]],
    shortlist_rows: Sequence[dict[str, str]],
    excluded_rows: Sequence[dict[str, str]],
    valuations: Mapping[str, dict[str, str]],
    valuation_evidence_rows: Sequence[dict[str, str]],
    trademark_summary_rows: Sequence[dict[str, str]],
    trademark_evidence_rows: Sequence[dict[str, str]],
    sample_rows: Sequence[dict[str, str]],
) -> dict[str, object]:
    rdap_counts = Counter(row["classification"] for row in rdap_latest.values())
    whois_counts = Counter(row["classification"] for row in whois_latest.values())
    valuation_types = Counter(row["evidence_type"] for row in valuation_evidence_rows)
    summary_counts = Counter(row["status"] for row in trademark_summary_rows)
    sourcewide = Counter(
        row["source"] for row in trademark_evidence_rows if row.get("domain", "").strip() == "*"
    )
    return {
        "candidates": len(candidate_rows),
        "candidate_categories": dict(Counter(row["category"] for row in candidate_rows)),
        "rdap_latest": dict(rdap_counts),
        "whois_latest": dict(whois_counts),
        "scores": len(scores),
        "shortlist": len(shortlist_rows),
        "excluded": len(excluded_rows),
        "outside_shortlist_cap": max(
            0, len(candidate_rows) - len(shortlist_rows) - len(excluded_rows)
        ),
        "valuations": len(valuations),
        "unsupported_valuations": sum(
            row["status"] == "insufficiently supported" for row in valuations.values()
        ),
        "valuation_evidence": len(valuation_evidence_rows),
        "valuation_evidence_types": dict(valuation_types),
        "trademark_summary": len(trademark_summary_rows),
        "trademark_summary_status": dict(summary_counts),
        "trademark_evidence": len(trademark_evidence_rows),
        "trademark_sourcewide_rows": dict(sourcewide),
        "review_sample": len(sample_rows),
    }


def _methodology_rows(
    metrics: Mapping[str, object],
    paths: Mapping[str, Path],
    hashes: Mapping[str, str],
    digest: str,
    tables: Mapping[str, CSVTable],
    rdap_latest: Mapping[str, dict[str, str]],
    whois_latest: Mapping[str, dict[str, str]],
    reserved: Mapping[str, dict[str, str]],
    restricted: Mapping[str, dict[str, str]],
    screening_counts: Mapping[str, Mapping[str, int]],
    unresolved: bool,
) -> tuple[tuple[object, ...], ...]:
    rdap_latest_counts = Counter(row["classification"] for row in rdap_latest.values())
    whois_latest_counts = Counter(row["classification"] for row in whois_latest.values())
    weights = "; ".join(f"{name} {weight * 100:g}%" for name, weight in SCORE_WEIGHTS.items())
    coverage = "; ".join(
        f"{source}: completed {screening_counts[source]['completed']}/{min(100, int(metrics['shortlist']))}, "
        f"source-wide blocked {screening_counts[source]['sourcewide_blocked']}, "
        f"query blocked {screening_counts[source]['query_blocked']}, "
        f"unverified/missing {screening_counts[source]['unverified']}"
        for source in REQUIRED_SOURCES
    )
    resume_requests: dict[tuple[str, str, str], list[dict[str, str]]] = {}
    for evidence_row in tables["trademark_evidence"].rows:
        evidence_id = evidence_row.get("evidence_id", "")
        if not evidence_id.startswith("RESUME-USPTO-"):
            continue
        note = evidence_row.get("notes", "")
        is_control = (
            evidence_row.get("domain") == "*"
            and evidence_row.get("query", "").casefold() == "google"
            and "control query" in note.casefold()
        )
        if evidence_row.get("domain") == "*" and not is_control:
            continue
        if not is_control and any(
            marker in note.casefold()
            for marker in ("no per-name query was sent", "query not sent", "no request made")
        ):
            continue
        if not is_control and not (
            evidence_row.get("proof_kind") == "api_query_summary"
            or "http 200 json response was received" in note.casefold()
            or "stopped on this request" in note.casefold()
            or "transient failure after one allowed retry" in note.casefold()
        ):
            continue
        key = (
            evidence_row.get("domain", ""),
            evidence_row.get("query", ""),
            evidence_row.get("utc_timestamp", ""),
        )
        resume_requests.setdefault(key, []).append(evidence_row)

    def _request_attempts(rows: list[dict[str, str]]) -> int:
        attempts = [1]
        for evidence_row in rows:
            for token in evidence_row.get("notes", "").replace(";", " ").split():
                if token.startswith("attempts="):
                    try:
                        attempts.append(int(token.partition("=")[2]))
                    except ValueError:
                        pass
        return max(attempts)

    resume_posts = sum(_request_attempts(rows) for rows in resume_requests.values())
    resume_controls = sum(domain == "*" for domain, _query, _timestamp in resume_requests)
    resume_candidate_requests = len(resume_requests) - resume_controls
    resume_http_200 = sum(
        any("http 200 json" in row.get("notes", "").casefold() for row in rows)
        for rows in resume_requests.values()
    )
    resume_http_403 = sum(
        sum(row.get("notes", "").casefold().count("http 403") for row in rows)
        for rows in resume_requests.values()
    )
    resume_retries = sum(max(0, _request_attempts(rows) - 1) for rows in resume_requests.values())
    resume_unverified_captures = sum(
        any("local evidence write failed" in row.get("notes", "").casefold() for row in rows)
        for rows in resume_requests.values()
    )
    control_label = "google control" if resume_controls == 1 else "google controls"
    resume_outcomes = (
        f"PKG-G and T-18 rework recorded {resume_posts} USPTO POST attempts: {resume_controls} {control_label} and "
        f"{resume_candidate_requests} candidate queries; {resume_http_200} returned HTTP 200 JSON, "
        f"{resume_http_403} returned HTTP 403, and {resume_retries} retries were used. "
        f"{resume_unverified_captures} historical response row records a local evidence-write failure; latest query evidence wins."
    )
    final_top_domains = {
        row["domain"].casefold() for row in tables["shortlist"].rows[:100]
    }
    final_flags = sorted(
        row["domain"].removesuffix(".si")
        for row in tables["trademark_summary"].rows
        if row["domain"].casefold() in final_top_domains and row["status"] == "needs legal review"
    )
    sales = metrics["valuation_evidence_types"]
    val_summary = ", ".join(
        f"{sales.get(kind, 0)} {kind.replace('_', ' ')}"
        for kind in ("completed_sale", "asking_price", "unsupported")
    )
    source_urls = sorted(
        {rules.GENERAL_TERMS_URL, rules.TARIFF_URL}
        | {
            row.get("source_url", "")
            for table in (tables["reserved_names"], tables["restricted_names"])
            for row in table.rows
            if row.get("source_url", "")
        }
    )
    rows: list[tuple[object, ...]] = []

    def add(topic: str, detail: str) -> None:
        rows.append((topic, detail, "", ""))

    add(
        "Summary",
        f"{metrics['shortlist']} ranked opportunities from {metrics['candidates']} unique candidate rows. "
        f"Trademark coverage is {'BLOCKED' if unresolved else 'screened for required query coverage'}; "
        f"all {metrics['unsupported_valuations']} valuations are insufficiently supported.",
    )
    add("Snapshot ID", f"sha256:{digest}")
    add(
        "Observation time",
        "RDAP, WHOIS, trademark, and sale dates are copied exactly from their source rows. "
        "For append-only RDAP/WHOIS files, the current result is the last row for each domain.",
    )
    add(
        "Export time",
        "The workbook does not store a wall-clock build time. The source-manifest SHA-256 is the deterministic snapshot identifier, so identical source paths and bytes yield identical cell content and styles.",
    )
    add("RDAP history timestamp range", _timestamp_range(tables["rdap_evidence"].rows, "utc_timestamp"))
    add("WHOIS history timestamp range", _timestamp_range(tables["whois_evidence"].rows, "utc_timestamp"))
    add("Trademark evidence timestamp range", _timestamp_range(tables["trademark_evidence"].rows, "utc_timestamp"))
    add("Latest RDAP timestamp range", _timestamp_range(list(rdap_latest.values()), "utc_timestamp"))
    add("Latest WHOIS timestamp range", _timestamp_range(list(whois_latest.values()), "utc_timestamp"))
    add(
        "Candidate categories",
        "; ".join(f"{key}: {value}" for key, value in sorted(metrics["candidate_categories"].items())),
    )
    add(
        "Category availability table columns",
        "Category | candidate pool | eligible (RDAP 404 + WHOIS not_registered, known rules pass) | final shortlist | median global eligible rank | minimum global eligible rank",
    )
    candidate_rows = tables["candidates"].rows
    candidate_categories = Counter(row["category"] for row in candidate_rows)
    selected_categories = Counter(row["category"] for row in tables["shortlist"].rows)
    score_by_domain = {row["domain"].casefold(): row for row in tables["scores"].rows}
    eligible_rows: list[tuple[str, str, dict[str, str]]] = []
    eligible_for_selection: list[tuple[str, dict[str, str], dict[str, str], object]] = []
    for candidate in candidate_rows:
        label = candidate["domain_label"].strip().casefold()
        domain = f"{label}.si"
        rdap = rdap_latest.get(domain)
        whois = whois_latest.get(domain)
        if (
            rules.classify(label) == "ok"
            and rdap
            and rdap["classification"] == "potentially_unregistered"
            and whois
            and whois["classification"] == "not_registered"
            and domain in score_by_domain
        ):
            eligible_rows.append((domain, candidate["category"], score_by_domain[domain]))
            eligible_for_selection.append(
                (domain, candidate, score_by_domain[domain], candidate)
            )
    eligible_rows.sort(
        key=lambda item: (
            -Decimal(item[2]["total"]),
            len(item[0].removesuffix(".si")),
            item[0],
        )
    )
    eligible_ranks: dict[str, list[int]] = {category: [] for category in candidate_categories}
    for global_rank, (_domain_name, category, _score) in enumerate(eligible_rows, start=1):
        eligible_ranks[category].append(global_rank)
    for category in sorted(candidate_categories):
        ranks = eligible_ranks[category]
        add(
            f"Availability: {category}",
            f"{category} | {candidate_categories[category]} | {len(ranks)} | "
            f"{selected_categories.get(category, 0)} | {median(ranks) if ranks else ''} | "
            f"{min(ranks) if ranks else ''}",
        )
    _selected_for_methodology, family_reasons = select_eligible(eligible_for_selection)
    family_cap_count = sum(reason == "family cap" for reason in family_reasons.values())
    add(
        "Latest RDAP counts",
        "; ".join(f"{key}: {rdap_latest_counts.get(key, 0)}" for key in ("registered", "potentially_unregistered", "unverified")),
    )
    add(
        "Latest WHOIS counts",
        "; ".join(f"{key}: {whois_latest_counts.get(key, 0)}" for key in ("not_registered", "registered", "unverified")),
    )
    add("Unverified sheet", "Zero rows means no latest RDAP result is unverified or missing and no necessary WHOIS second check is pending; trademark uncertainty remains shown separately.")
    add("Scores", f"{metrics['scores']} RDAP-404 rows; no score is fabricated for names outside the score artifact.")
    add("Ranking weights", weights)
    add("Score convention", f"Total is a 0-100 heuristic; higher is stronger. Components are integer 0-100; total uses Decimal ROUND_HALF_EVEN to one decimal. Trademark scores: {TRADEMARK_CLEAR_SCORE} complete limited screen/no obvious match; {TRADEMARK_CONFLICT_SCORE} possible conflict; {TRADEMARK_UNVERIFIED_SCORE} incomplete, blocked, unverified, or not yet screened. Unknown risk is penalized without weight renormalization; clear screening can add up to 4.5 weighted points over an unscreened name. No score implies clearance.")
    add("Brandability factors", "Label length 25%, pronounceability 22%, spelling clarity 15%, candidate quality 15%, letter-shape score 10%, vowel balance 4%, transition balance 3%, distinct letters 3%, onset/ending clarity 3%; reviewed phrase-hack adjustments and exact negative-label penalty apply.")
    add("Commercial scoring factors", "Tier bands: 1=85-100, 2=65-85, 3=45-65; no-root/name-like=30-50; dictionary words without industry roots=25-45. Compound demand is strongest root x 0.88 + second root x 0.12, plus 4 when both are listed commercial roots. Rootless coinages use a 34 base and per-label shape/quality adjustments; exact reviewed negative labels lose 50. See the complete demand-root table below.")
    add("Resale heuristic factors", "Per label: length, syllables, phonetic ease, shape, distinct letters, edge clarity, root demand/breadth, exact-word commonness, memorability, hack rarity, and ambiguity/repetition penalties. Rootless invented/name-like labels cap at 65; brandability caps at 65. Strong roots and broad-buyer common words can exceed that cap. This is not a price estimate.")
    add("Shortness and hack factors", "Shortness uses length, syllables, pronunciation/shape, letter variety and edge clarity. Short names and readable hacks get within-component bonuses. All six eligible CVCV shorts are selected at ranks 14, 15, 16, 18, 19, and 28; no three-letter short name was eligible. Three of 15 eligible hacks rank 1-3; hack selection remains quality constrained.")
    add("Rationale specificity", "Each shortlist rationale is <=140 ASCII characters and names a concrete fact: matched root and buyer pool, compound head/tail, common-word sense, hack completion, or label ending/letter shape.")
    add("Negative-label penalties", ", ".join(sorted(_negative_connotations())))
    dictionary_labels = sorted(
        row["domain"].removesuffix(".si")
        for row in tables["shortlist"].rows
        if row["category"] == "dictionary"
    )
    add("Dictionary and given-name classification", f"The exact curated data/given_names.txt list moves reviewed given names and pseudo-names to invented with tag given_name; endings alone are not sufficient evidence. Final dictionary rows ({len(dictionary_labels)}): {', '.join(dictionary_labels)}. Candidate labels remain in the 6,204-row pool.")
    add("Additional reviewed penalties", "Added after the final-label review: bloodmonitor, bloodtrack, gynclinic, hemoclinic, tumormarker. The nine named negative/medical/pejorative terms are absent from the final 500. Penalties change scores, not pool membership.")
    for root, (tier, demand, breadth, buyer_summary) in sorted(
        COMMERCIAL_ROOTS.items(), key=lambda item: (item[1][0], item[0])
    ):
        add(
            f"Demand root: {root}",
            f"tier={tier}; demand={demand}/100; buyer breadth={breadth}; buyer pool={buyer_summary}",
        )
    add("Eligibility", "A shortlist row requires known rules to pass, latest RDAP potentially_unregistered (HTTP 404), and latest WHOIS not_registered. WHOIS is a separate channel of the same registry, not another registrar/operator.")
    top100_counts = Counter(row["category"] for row in tables["shortlist"].rows[:100])
    top100_noninvented = sum(
        row["category"] != "invented" for row in tables["shortlist"].rows[:100]
    )
    add("Ranking and ties", "Total descending, label length ascending, then case-folded domain ascending. Selection has a 500-row limit, 100 per category (20% of the shortlist), and a 35-per-category top-100 limit. Family caps: at most 6 invented/given-name labels per final-three-letter ending and per first-four-letter prefix; at most 10 compound/blend labels per head or tail. Overflow remains eligible and is shown in All candidates as family cap. These limits diversify selection without removing pool labels.")
    add("Family-cap exclusions", f"{family_cap_count} eligible labels were not selected because an invented/given-name or compound/blend family cap was reached; they remain in All candidates with exclusion reason 'family cap'.")
    add("Top-100 diversity", f"Category counts: {'; '.join(f'{key} {value}' for key, value in sorted(top100_counts.items()))}. Non-invented keyword/real-word/short/hack rows: {top100_noninvented}/100. Maximum category count: {max(top100_counts.values(), default=0)}/35.")
    score_tuples = {
        tuple(row[name] for name in ("brandability", "commercial", "shortness", "hack", "resale"))
        for row in tables["shortlist"].rows
    }
    commercial_scores = [int(row["commercial"]) for row in tables["shortlist"].rows]
    industry_scores = [
        int(row["commercial"])
        for row in tables["shortlist"].rows
        if row["category"] in {"ai", "technology", "business_finance"}
    ]
    invented_scores = [
        int(row["commercial"])
        for row in tables["shortlist"].rows
        if row["category"] == "invented"
    ]
    commercial_spread = f"{pstdev(commercial_scores):.2f}" if commercial_scores else "n/a"
    industry_mean = f"{mean(industry_scores):.2f}" if industry_scores else "n/a"
    invented_mean = f"{mean(invented_scores):.2f}" if invented_scores else "n/a"
    add("Score diversity metrics", f"Final {len(tables['shortlist'].rows)}: {len(score_tuples)} distinct five-component tuples and {len({row['rationale'] for row in tables['shortlist'].rows})} distinct rationales. Commercial population standard deviation={commercial_spread}; AI/technology/business_finance mean={industry_mean}; invented mean={invented_mean}.")
    for category in sorted({row["category"] for row in tables["shortlist"].rows}):
        category_rows = [
            row for row in tables["shortlist"].rows if row["category"] == category
        ]
        add(
            f"Score means: {category}",
            f"rows={len(category_rows)}; commercial={mean(int(row['commercial']) for row in category_rows):.2f}; resale={mean(int(row['resale']) for row in category_rows):.2f}",
        )
    add("Exclusions", f"{metrics['excluded']} ineligible rows are in the exclusion artifact; {metrics['outside_shortlist_cap']} eligible rows fall outside the final ranked selection.")
    add("Trademark query coverage", coverage)
    add("Trademark status", "PKG-G screened the 44 previously USPTO-blocked names and completed three entrant rounds. The T-18 rework fixed ASCII evidence sanitization, recovered powdery's fuzzy query, queried denimwear exact and near, and screened soillab as the replacement entrant while USPTO remained live. Relative to the pre-rework PKG-G top 100, soillab entered and denimwear exited; denimwear's conflict evidence remains recorded. Final top-100 USPTO coverage: " + coverage + ". Relative to T-17, seven names entered (examhub, bankingvault, roboticlink, roboticloop, biomasshub, roboticflow, biomasslab) and seven exited (robotsystem, gyration, crophub, croplab, agrihub, cryptoworks, parcelshop). Names leaving retain their evidence. USPTO conflicts take precedence over incomplete other sources and score 25; otherwise incomplete coverage scores 40.")
    add("Manual trademark checks the owner can do", "Official pages: TMview https://www.tmdn.org/tmview/; UKIPO https://trademarks.ipo.gov.uk/ipo-tmtext; USPTO https://tmsearch.uspto.gov/. TMview source-wide refused: connection closed (evidence REWORK1-TMVIEW-PROBES, 2026-10-13T18:02:42Z); per-name results are unverified because of that refusal. UKIPO source-wide refused: automated access refused (evidence PKGD-0003, 2026-10-13T17:09:52Z); per-name results are blocked because of that refusal. No further TMview or UKIPO attempts were made. USPTO's prior HTTP 403 at 2026-10-13T18:05:11Z is retained (REWORK1-USPTO-00259; REWORK1-USPTO-SOURCE-STOP). PKG-G originally made 117 USPTO POSTs (one control and 116 candidates); T-18 added six (one control and five candidates). The combined total is " + str(resume_posts) + " sequential requests, at least five seconds apart; all " + str(resume_http_200) + " returned HTTP 200 JSON, with no new 403, CAPTCHA, login, non-JSON response, timeout, 5xx, or retry. No result is a clearance.")
    add("Completion-mode research validation", "Use python -m si_domains validate-research --candidates data/candidates.csv --completion-mode for an offline record check. Each top-100 name/source pair needs completed exact and near queries, or dated blocked/unverified rows with an official URL and an evidence-backed reason (HTTP status, CAPTCHA, connection closed, source-stop ID, local evidence-write failure, or the documented three-round limit). Missing pairs fail; incomplete sources score 40, conflicts score 25, and 85 requires all three sources complete with no match. Every top-100 opportunity must have a trademark status other than not yet screened. Strict mode without the flag still exits 2 until required queries complete.")
    add("Trademark mapping rationale", f"One source of truth: limited_screen_no_obvious_match={TRADEMARK_CLEAR_SCORE}; needs legal review={TRADEMARK_CONFLICT_SCORE}; blocked/unverified/not yet screened={TRADEMARK_UNVERIFIED_SCORE}. The documented {TRADEMARK_UNVERIFIED_SCORE} is below clear and above conflict; no renormalization is used. Unknown risk loses up to 4.5 weighted points versus a clear result; a finding loses more.")
    add("Trademark similarity rule", "Flag legal review for identical word marks, whole-word containment, edit distance <=1 for labels of 5+ characters, or identical stem, in any class. Classes 9, 35, 36, 38, 41, and 42 are higher relevance. Uncertain is never clear.")
    add("Trademark probe outcome", "Before PKG-G, USPTO history recorded one HTTP 403 on bankingdesk; it remains in the evidence log. " + resume_outcomes + " The T-18 requests were: one google control; powdery [fuzzy-1] (2097 total, 10 returned, 10 similar); denimwear exact and fuzzy-1 (2 returned and similar per query); and soillab exact (0 hits) and fuzzy-1 (12 total, 10 returned, 10 similar). All five candidate queries and the control returned HTTP 200 JSON. Final top-100 source coverage: " + coverage + ". Current names needing legal review: " + (", ".join(final_flags) if final_flags else "none") + ". Thirteen newly flagged names are autokernel, autonetwork, autostream, carboncell, cashops, jewelshop, keyscan, pharmascan, powdery, reveler, robotspark, soilhub, and soillab.")
    add("Trademark limitations", "Exact public-query coverage is not comprehensive phonetic, class, national-office, or worldwide clearance. Possible conflicts need professional legal review; unknown risk scores 40, below a complete no-obvious-match screen at 85 and above a conflict at 25.")
    add("Valuation evidence", f"Evidence records: {val_summary}. No current evidence meets the direct-comparable rule for any shortlist row; unsupported valuations have blank currency and low/high cells.")
    add("Availability limits", "RDAP and WHOIS are timestamped registry-channel observations, not registrar checkout confirmation, eligibility confirmation, or a guarantee of successful registration.")
    add("Pricing limits", "Public registrar confirmation and registrar retail/premium pricing are unknown. The rule inventory has no premium-name list; an empty premium list is not evidence of ordinary pricing.")
    add("Registry rule inventory", f"{len(reserved)} reserved and {len(restricted)} restricted names are evidence-backed. Known-list limits and classification scope are described in docs/si_registry_rules.md and docs/special_classes.md.")
    add("Rule source URLs", "; ".join(source_urls) if source_urls else "No reserved/restricted inventory rows.")
    add("Special-sheet alias", "Excel sheet name `Special classes` is the documented alias for `Reserved/Restricted/Invalid/Premium` because Excel forbids `/` and limits names to 31 characters. Scope marks candidate rows separately from rule-reference inventory rows.")
    add("Provenance notes", "See docs/si_registry_rules.md, docs/special_classes.md, output/research_methodology.md, and output/valuation_sources_notes.md.")
    add("Research exit", "Strict validate-research exits 2 after structure passes because exact-plus-near coverage is incomplete: " + coverage + ". Blocked/unverified rows do not complete a search. --completion-mode checks outcome-record completeness and status/score honesty without network requests; it does not establish clearance. Delivery validation is separate.")
    for name in _INPUT_ORDER:
        path = paths[name]
        rows.append((f"Source CSV: {name}", "Byte-level SHA-256 of the exact input file.", _display_path(path), hashes[name]))
    return tuple(rows)


def _sheet_data(
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    *,
    freeze_panes: str | None = "A2",
    wrap_columns: Sequence[str] = (),
    integer_columns: Sequence[str] = (),
    decimal_columns: Sequence[str] = (),
    filters: bool = True,
) -> SheetData:
    return SheetData(
        tuple(headers),
        tuple(tuple(row) for row in rows),
        freeze_panes,
        filters,
        frozenset(wrap_columns),
        frozenset(integer_columns),
        frozenset(decimal_columns),
    )


def build_workbook_model(paths: Mapping[str, str | Path] | None = None) -> WorkbookModel:
    """Validate the snapshot and construct the deterministic workbook cell model."""

    resolved = _resolved_inputs(paths)
    tables = _read_inputs(resolved)
    candidate_rows = list(tables["candidates"].rows)
    candidates = _index_unique(
        candidate_rows, "domain_label", "candidate", normalize_domain=True
    )
    # Candidate-map keys are full domains, not labels.
    candidates = {f"{label}.si": row for label, row in candidates.items()}
    candidate_domains = set(candidates)
    if any(not row["category"].strip() for row in candidate_rows):
        raise ValueError("candidate category cannot be blank")

    reserved, restricted = rules.load_name_rules(
        resolved["reserved_names"], resolved["restricted_names"]
    )
    rdap_latest, whois_latest = _validate_registry_rows(
        candidates, tables["rdap_evidence"].rows, tables["whois_evidence"].rows
    )

    # Recompute current summary before validating component scores.
    valuation_evidence_rows, _valuation_index = read_valuation_evidence(
        resolved["valuation_evidence"]
    )
    trademark_evidence_rows, _trademark_index = read_trademark_evidence(
        resolved["trademark_evidence"]
    )
    candidate_context, query_manifest = read_candidate_context(resolved["candidates"])
    if set(candidate_context) != candidate_domains:
        raise ValueError("candidate metadata does not resolve to the exact candidate domain set")
    summary_map = read_trademark_summary(
        resolved["trademark_summary"],
        evidence_path=resolved["trademark_evidence"],
        candidate_domains=candidate_domains,
        candidates_path=resolved["candidates"],
    )
    scores = _score_index(
        tables["scores"].rows, candidates, rdap_latest, summary_map
    )
    shortlist, excluded, shortlisted_domains, selection_reasons = _validate_shortlist_and_excluded(
        candidate_rows,
        candidates,
        rdap_latest,
        whois_latest,
        scores,
        tables["shortlist"].rows,
        tables["excluded"].rows,
        reserved,
        restricted,
    )
    valuations, opportunities, unresolved, screening_counts = _validate_research_tables(
        tables=tables,
        candidates=candidates,
        candidate_rows=candidate_rows,
        rdap_latest=rdap_latest,
        whois_latest=whois_latest,
        scores=scores,
        shortlist=shortlist,
        shortlisted_domains=shortlisted_domains,
        reserved=reserved,
        restricted=restricted,
    )

    unverified_rows: list[dict[str, object]] = []
    all_candidate_rows: list[dict[str, object]] = []
    special_rows: list[dict[str, object]] = []
    candidate_header_collisions = set(tables["candidates"].headers) & {
        "domain", "latest_rdap_http_status", "latest_rdap_classification", "latest_rdap_timestamp",
        "latest_rdap_attempts", "latest_rdap_error", "latest_whois_classification",
        "latest_whois_timestamp", "latest_whois_matched_pattern", "known_rule_class",
        "known_rule_id", "known_rule_source_url", "known_rule_explanation", "registry_status",
        "exclusion_reason", "shortlist_rank", *_SCORE_COMPONENTS, "score_total",
    }
    if candidate_header_collisions:
        raise ValueError(f"candidate columns collide with workbook joins: {sorted(candidate_header_collisions)[0]}")

    selected = set(shortlist)
    excluded_reason = {domain: row["reason"] for domain, row in excluded.items()}
    for candidate in candidate_rows:
        label_original = candidate["domain_label"]
        label = label_original.strip().casefold()
        domain = f"{label}.si"
        rdap_row = rdap_latest.get(domain)
        whois_row = whois_latest.get(domain)
        rule_class = _rule_class(label, reserved, restricted)
        classification, rule_id, rule_url = rules.explain(
            label, reserved_rules=reserved, restricted_rules=restricted
        )
        explanation = ""
        if classification != "ok":
            rule_id, rule_url, explanation = _rule_explanation(
                label, classification, reserved, restricted
            )
            special_rows.append(
                {
                    "scope": "candidate",
                    "domain": domain,
                    "class": classification,
                    "rule_id": rule_id,
                    "source_url": rule_url,
                    "explanation": explanation,
                }
            )

        reason = excluded_reason.get(domain) or selection_reasons.get(domain)
        if domain in selected:
            reason = ""
        elif reason is None and rule_class == "ok" and _reason_for(label, rule_class, rdap_row, whois_row) is None:
            reason = "outside_shortlist_cap"
        candidate_status = _candidate_status(
            domain, rule_class, rdap_row, whois_row, selected
        )
        score = scores.get(domain)
        all_candidate_rows.append(
            {
                "domain": domain,
                **candidate,
                "latest_rdap_http_status": int(rdap_row["http_status"]) if rdap_row and rdap_row["http_status"] else None,
                "latest_rdap_classification": rdap_row["classification"] if rdap_row else "missing",
                "latest_rdap_timestamp": rdap_row["utc_timestamp"] if rdap_row else "",
                "latest_rdap_attempts": int(rdap_row["attempts"]) if rdap_row else None,
                "latest_rdap_error": rdap_row.get("error", "") if rdap_row else "",
                "latest_whois_classification": whois_row["classification"] if whois_row else "missing",
                "latest_whois_timestamp": whois_row["utc_timestamp"] if whois_row else "",
                "latest_whois_matched_pattern": whois_row.get("matched_pattern", "") if whois_row else "",
                "known_rule_class": rule_class,
                "known_rule_id": rule_id,
                "known_rule_source_url": rule_url,
                "known_rule_explanation": explanation or "Passes the known local syntax and name lists; this does not establish availability.",
                **{name: int(score[name]) if score else None for name in _SCORE_COMPONENTS},
                "score_total": _excel_numeric(score["total"], "total", domain) if score else None,
                "shortlist_rank": int(shortlist[domain]["rank"]) if domain in selected else None,
                "registry_status": candidate_status,
                "exclusion_reason": reason or "",
            }
        )

        if rule_class == "ok":
            unverified_reason = ""
            if rdap_row is None:
                unverified_reason = "latest RDAP evidence is missing"
            elif rdap_row["classification"] == "unverified":
                unverified_reason = "latest RDAP result is unverified"
            elif rdap_row["classification"] == "potentially_unregistered":
                if whois_row is None:
                    unverified_reason = "required WHOIS second check is missing"
                elif whois_row["classification"] == "unverified":
                    unverified_reason = "required WHOIS second check is unverified"
            if unverified_reason:
                unverified_rows.append(
                    {
                        "domain": domain,
                        "category": candidate["category"],
                        "reason": unverified_reason,
                        "rdap_http_status": int(rdap_row["http_status"]) if rdap_row and rdap_row["http_status"] else None,
                        "rdap_classification": rdap_row["classification"] if rdap_row else "missing",
                        "rdap_timestamp": rdap_row["utc_timestamp"] if rdap_row else "",
                        "rdap_attempts": int(rdap_row["attempts"]) if rdap_row else None,
                        "rdap_error": rdap_row.get("error", "") if rdap_row else "",
                        "whois_classification": whois_row["classification"] if whois_row else "missing",
                        "whois_timestamp": whois_row["utc_timestamp"] if whois_row else "",
                        "whois_matched_pattern": whois_row.get("matched_pattern", "") if whois_row else "",
                        "whois_error": whois_row.get("error", "") if whois_row else "",
                    }
                )

    # Rule-reference entries stay separate from queried in-pool candidates.
    for rule_class, entries in (("reserved", reserved), ("restricted", restricted)):
        for label, entry in entries.items():
            classification, rule_id, source_url = rules.explain(
                label, reserved_rules=reserved, restricted_rules=restricted
            )
            if classification != rule_class or not entry.get("note", "").strip():
                raise ValueError(f"invalid {rule_class} rule inventory entry: {label}")
            special_rows.append(
                {
                    "scope": "rule_reference",
                    "domain": f"{label}.si",
                    "class": rule_class,
                    "rule_id": rule_id,
                    "source_url": source_url,
                    "explanation": entry["note"],
                }
            )

    candidate_output_headers = (
        "domain",
        *tables["candidates"].headers,
        "latest_rdap_http_status",
        "latest_rdap_classification",
        "latest_rdap_timestamp",
        "latest_rdap_attempts",
        "latest_rdap_error",
        "latest_whois_classification",
        "latest_whois_timestamp",
        "latest_whois_matched_pattern",
        "known_rule_class",
        "known_rule_id",
        "known_rule_source_url",
        "known_rule_explanation",
        *_SCORE_COMPONENTS,
        "score_total",
        "shortlist_rank",
        "registry_status",
        "exclusion_reason",
    )
    candidate_integer_columns = frozenset(
        {"latest_rdap_http_status", "latest_rdap_attempts", "shortlist_rank", *_SCORE_COMPONENTS}
    )
    candidate_decimal_columns = frozenset({"score_total"})
    all_candidate_values = [tuple(row.get(header, "") for header in candidate_output_headers) for row in all_candidate_rows]

    shortlist_headers = _SHORTLIST_EXPORT_FIELDS
    shortlist_output_rows: list[tuple[object, ...]] = []
    for opportunity in sorted(tables["opportunities"].rows, key=lambda row: int(row["rank"])):
        domain = _domain(opportunity["domain"], "opportunities")
        rdap_row = rdap_latest.get(domain)
        whois_row = whois_latest.get(domain)
        if domain not in shortlist or rdap_row is None or whois_row is None:
            raise ValueError(f"shortlist opportunity lacks current registry evidence for {domain}")
        values: dict[str, object] = {
            field: opportunity.get(field, "") for field in OPPORTUNITY_FIELDS
        }
        for field in ("rank", *_SCORE_COMPONENTS):
            values[field] = _excel_numeric(str(values[field]), field, domain)
        values["total"] = _excel_numeric(str(values["total"]), "total", domain)
        for field in ("valuation_low", "valuation_high"):
            values[field] = _excel_numeric(str(values[field]), field, domain)
        values["rdap_http_status"] = int(rdap_row["http_status"]) if rdap_row["http_status"] else None
        values["rdap_classification"] = rdap_row["classification"]
        values["whois_classification"] = whois_row["classification"]
        shortlist_output_rows.append(tuple(values[field] for field in shortlist_headers))

    special_headers = ("scope", "domain", "class", "rule_id", "source_url", "explanation")
    special_rows.sort(key=lambda row: (row["scope"], row["class"], row["domain"]))
    special_values = [tuple(row[field] for field in special_headers) for row in special_rows]
    unverified_headers = (
        "domain", "category", "reason", "rdap_http_status", "rdap_classification", "rdap_timestamp",
        "rdap_attempts", "rdap_error", "whois_classification", "whois_timestamp", "whois_matched_pattern", "whois_error",
    )
    unverified_values = [tuple(row[field] for field in unverified_headers) for row in unverified_rows]

    def raw_rows(table_name: str) -> tuple[tuple[object, ...], ...]:
        return tuple(tuple(row.get(header, "") for header in tables[table_name].headers) for row in tables[table_name].rows)

    hashes = {name: _sha256(resolved[name]) for name in _INPUT_ORDER}
    manifest = "\n".join(
        f"{name}\t{_display_path(resolved[name])}\t{hashes[name]}" for name in _INPUT_ORDER
    ).encode("utf-8")
    digest = hashlib.sha256(manifest).hexdigest()
    metrics = _metric_counts(
        candidate_rows,
        rdap_latest,
        whois_latest,
        scores,
        tables["shortlist"].rows,
        tables["excluded"].rows,
        valuations,
        valuation_evidence_rows,
        tables["trademark_summary"].rows,
        trademark_evidence_rows,
        tables["sample"].rows,
    )
    methodology_values = _methodology_rows(
        metrics,
        resolved,
        hashes,
        digest,
        tables,
        rdap_latest,
        whois_latest,
        reserved,
        restricted,
        screening_counts,
        unresolved,
    )

    score_rows = tables["scores"].rows
    score_output_rows = []
    for row in score_rows:
        domain = _domain(row["domain"], "scores")
        score_output_rows.append(
            tuple(
                _excel_numeric(row[field], field, domain)
                if field in {*_SCORE_COMPONENTS, "total"}
                else row.get(field, "")
                for field in tables["scores"].headers
            )
        )
    valuation_evidence_data = []
    for row in tables["valuation_evidence"].rows:
        valuation_evidence_data.append(
            tuple(
                _excel_numeric(row[field], field, row.get("evidence_id", ""))
                if field == "amount" and row.get("amount", "")
                else row.get(field, "")
                for field in tables["valuation_evidence"].headers
            )
        )
    rdap_evidence_data = []
    for row in tables["rdap_evidence"].rows:
        rdap_evidence_data.append(
            tuple(
                int(row[field]) if field in {"http_status", "attempts"} and row.get(field, "") else row.get(field, "")
                for field in tables["rdap_evidence"].headers
            )
        )
    sample_data = tuple(
        tuple(
            int(row[field]) if field == "rank" and row.get(field, "") else row.get(field, "")
            for field in tables["sample"].headers
        )
        for row in tables["sample"].rows
    )

    sheets: dict[str, SheetData] = {
        "Shortlist": _sheet_data(
            shortlist_headers,
            shortlist_output_rows,
            freeze_panes="F2",
            wrap_columns=("rationale", "valuation_rationale", "trademark_coverage", "trademark_evidence_ids"),
            integer_columns=("rank", *_SCORE_COMPONENTS, "rdap_http_status"),
            decimal_columns=("total", "valuation_low", "valuation_high"),
        ),
        "All candidates": _sheet_data(
            candidate_output_headers,
            all_candidate_values,
            freeze_panes="C2",
            wrap_columns=("tags", "seed", "latest_rdap_error", "known_rule_explanation", "exclusion_reason"),
            integer_columns=candidate_integer_columns,
            decimal_columns=candidate_decimal_columns,
        ),
        "Unverified": _sheet_data(
            unverified_headers,
            unverified_values,
            freeze_panes="C2",
            wrap_columns=("reason", "rdap_error", "whois_error"),
            integer_columns=("rdap_http_status", "rdap_attempts"),
        ),
        "Special classes": _sheet_data(
            special_headers,
            special_values,
            freeze_panes="C2",
            wrap_columns=("source_url", "explanation"),
        ),
        "Methodology": _sheet_data(
            ("Topic", "Details", "Input path", "SHA-256"),
            methodology_values,
            freeze_panes="A2",
            filters=False,
            wrap_columns=("Topic", "Details", "Input path", "SHA-256"),
        ),
        "Scores": _sheet_data(
            tables["scores"].headers,
            score_output_rows,
            freeze_panes="C2",
            wrap_columns=("rationale",),
            integer_columns=_SCORE_COMPONENTS,
            decimal_columns=("total",),
        ),
        "RDAP evidence": _sheet_data(
            tables["rdap_evidence"].headers,
            rdap_evidence_data,
            freeze_panes="C2",
            wrap_columns=("error",),
            integer_columns=("http_status", "attempts"),
        ),
        "WHOIS evidence": _sheet_data(
            tables["whois_evidence"].headers,
            raw_rows("whois_evidence"),
            freeze_panes="C2",
            wrap_columns=("matched_pattern", "error"),
        ),
        "Valuation evidence": _sheet_data(
            tables["valuation_evidence"].headers,
            valuation_evidence_data,
            freeze_panes="C2",
            wrap_columns=("source_url", "relevance", "limitations"),
            decimal_columns=("amount",),
        ),
        "Trademark evidence": _sheet_data(
            tables["trademark_evidence"].headers,
            raw_rows("trademark_evidence"),
            freeze_panes="C2",
            wrap_columns=("source_url", "result_url", "mark_name", "notes", "proof_path"),
        ),
    }

    # Methodology has a source-metadata table; its strings are plain text too.
    return WorkbookModel(
        sheets=sheets,
        input_paths=resolved,
        input_hashes=hashes,
        snapshot_digest=digest,
        metrics=metrics,
        research_unresolved=unresolved,
        screening_counts=screening_counts,
    )


_HEADER_FILL = PatternFill("solid", fgColor="17365D")
_HEADER_FONT = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
_BODY_FONT = Font(name="Aptos", size=10, color="202020")
_KEY_FILL = PatternFill("solid", fgColor="DCE6F1")


def _cell_value(value: object) -> object:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        return value
    return value


def _excel_text_length(value: str) -> int:
    """Return Excel's UTF-16 cell length and reject lone surrogate code points."""

    return len(value.encode("utf-16-le")) // 2


def _validate_cell_text(model: WorkbookModel) -> None:
    """Fail before writing when a text cell cannot be preserved in XLSX/XML."""

    for sheet_name, data in model.sheets.items():
        for row_number, row in enumerate((data.headers, *data.rows), start=1):
            for column_number, value in enumerate(row, start=1):
                if not isinstance(value, str):
                    continue
                coordinate = f"{get_column_letter(column_number)}{row_number}"
                header = data.headers[column_number - 1] if row_number > 1 else "header"
                context = f"{sheet_name}!{coordinate} ({header})"
                try:
                    length = _excel_text_length(value)
                except UnicodeEncodeError as exc:
                    raise ValueError(f"invalid XML text in {context}: unpaired Unicode surrogate") from exc
                if length > 32767:
                    raise ValueError(
                        f"text in {context} exceeds Excel's 32767-character cell limit ({length} UTF-16 code units)"
                    )
                for character in value:
                    codepoint = ord(character)
                    if not (
                        codepoint in (0x9, 0xA, 0xD)
                        or 0x20 <= codepoint <= 0xD7FF
                        or 0xE000 <= codepoint <= 0xFFFD
                        or 0x10000 <= codepoint <= 0x10FFFF
                    ):
                        raise ValueError(
                            f"invalid XML character U+{codepoint:04X} in {context}"
                        )


def _estimated_wrapped_lines(value: object, width: float) -> int:
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").expandtabs(4)
    line_width = max(1, int(width) - 2)
    count = 0
    for explicit_line in text.split("\n"):
        wrapped = textwrap.wrap(
            explicit_line,
            width=line_width,
            break_long_words=True,
            break_on_hyphens=True,
            replace_whitespace=False,
            drop_whitespace=False,
            max_lines=_MAX_ESTIMATED_WRAP_LINES,
            placeholder="...",
        )
        count += max(1, len(wrapped))
        if count >= _MAX_ESTIMATED_WRAP_LINES:
            return _MAX_ESTIMATED_WRAP_LINES
    return max(1, count)


def _write_sheet(workbook: Workbook, name: str, data: SheetData) -> Worksheet:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = True
    widths: dict[str, float] = {}
    for column, header in enumerate(data.headers, start=1):
        values = [str(header)]
        values.extend(
            str(row[column - 1])
            for row in data.rows
            if row[column - 1] is not None
        )
        longest = max(
            (len(line.expandtabs(4)) for value in values for line in value.splitlines()),
            default=10,
        )
        width = min(max(longest + 2, 12), 50)
        if header in data.wrap_columns:
            width = min(max(width, 30), 60)
        if header in {"domain", "domain_label"}:
            width = max(width, 22)
        if header in {"rank", "shortlist_rank", "http_status", "rdap_http_status"}:
            width = max(10, min(width, 14))
        if name == "Methodology" and header == "Details":
            width = 90
        if name == "Methodology" and header == "Topic":
            width = 32
        widths[header] = width

    for column, header in enumerate(data.headers, start=1):
        cell = sheet.cell(1, column, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    header_index = {header: index + 1 for index, header in enumerate(data.headers)}
    for row_index, row in enumerate(data.rows, start=2):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column)
            safe_value = _cell_value(value)
            if isinstance(safe_value, str):
                cell.value = safe_value
                # Force a shared/inline string even when text starts with a formula prefix.
                cell.data_type = "s"
            else:
                cell.value = safe_value
            cell.font = _BODY_FONT
            header = data.headers[column - 1]
            cell.alignment = Alignment(
                vertical="top", wrap_text=header in data.wrap_columns
            )
            if header in data.integer_columns and safe_value is not None:
                cell.number_format = "0"
            elif header in data.decimal_columns and safe_value is not None:
                cell.number_format = "0.0"
            if name == "Shortlist" and header in {"rank", "domain", "total", "trademark", "trademark_status"}:
                cell.fill = _KEY_FILL
        if data.wrap_columns:
            required_lines = max(
                (
                    _estimated_wrapped_lines(value, widths[data.headers[column - 1]])
                    for column, value in enumerate(row, start=1)
                    if data.headers[column - 1] in data.wrap_columns and value is not None
                ),
                default=1,
            )
            sheet.row_dimensions[row_index].height = min(
                _MAX_EXCEL_ROW_HEIGHT, max(18, required_lines * 15)
            )

    if data.freeze_panes:
        sheet.freeze_panes = data.freeze_panes
    if data.filters:
        last_column = get_column_letter(max(1, len(data.headers)))
        sheet.auto_filter.ref = f"A1:{last_column}{max(1, len(data.rows) + 1)}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"

    for column, header in enumerate(data.headers, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = widths[header]
    return sheet


def create_workbook(model: WorkbookModel) -> Workbook:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for name, sheet_data in model.sheets.items():
        _write_sheet(workbook, name, sheet_data)
    workbook.properties.creator = "si_domains offline export"
    workbook.properties.title = "Ranked .si domain research snapshot"
    workbook.properties.subject = f"Snapshot sha256:{model.snapshot_digest}"
    deterministic_timestamp = datetime(2000, 1, 1, tzinfo=timezone.utc)
    workbook.properties.created = deterministic_timestamp.replace(tzinfo=None)
    workbook.properties.modified = deterministic_timestamp.replace(tzinfo=None)
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.calculation.calcMode = "manual"
    return workbook


def _reopen_in_all_modes(path: Path) -> None:
    if not zipfile.is_zipfile(path):
        raise ValueError("export is not a valid XLSX ZIP archive")
    for options in (
        {"read_only": False, "data_only": False},
        {"read_only": True, "data_only": False},
        {"read_only": False, "data_only": True},
    ):
        workbook = load_workbook(path, **options)
        workbook.close()


def export_workbook(
    model_or_paths: WorkbookModel | Mapping[str, str | Path] | None,
    output_path: str | Path,
) -> WorkbookModel:
    """Atomically write an XLSX; no source CSV is ever used as an output target."""

    model = (
        model_or_paths
        if isinstance(model_or_paths, WorkbookModel)
        else build_workbook_model(model_or_paths)
    )
    _validate_cell_text(model)
    output = Path(output_path).resolve()
    source_paths = {path.resolve() for path in model.input_paths.values()}
    if output in source_paths:
        raise ValueError(f"workbook output cannot overwrite a source CSV: {output}")
    if output.suffix.casefold() != ".xlsx":
        raise ValueError("workbook output path must end in .xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temp_name = tempfile.mkstemp(
        prefix=f".{output.stem}.", suffix=".xlsx", dir=output.parent
    )
    os.close(descriptor)
    temporary = Path(temp_name)
    try:
        workbook = create_workbook(model)
        workbook.save(temporary)
        workbook.close()
        _reopen_in_all_modes(temporary)
        os.replace(temporary, output)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return model
