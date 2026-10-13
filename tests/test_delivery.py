import csv
import shutil

import pytest
from openpyxl import load_workbook

from si_domains.cli import main
from si_domains.delivery import validate_delivery
from si_domains.trademarks import TRADEMARK_EVIDENCE_FIELDS, write_trademark_summary
from si_domains.workbook import export_workbook

from test_workbook import make_fixture


def _cli_inputs(paths):
    flags = {
        "candidates": "--candidates",
        "reserved_names": "--reserved-names",
        "restricted_names": "--restricted-names",
        "rdap_evidence": "--rdap-evidence",
        "whois_evidence": "--whois-evidence",
        "scores": "--scores",
        "shortlist": "--shortlist",
        "excluded": "--excluded",
        "opportunities": "--opportunities",
        "valuations": "--valuations",
        "valuation_evidence": "--valuation-evidence",
        "trademark_summary": "--trademark-summary",
        "trademark_evidence": "--trademark-evidence",
        "sample": "--sample",
    }
    return [item for name, flag in flags.items() for item in (flag, str(paths[name]))]


def _add_prior_screen_record(paths):
    evidence_path = paths["trademark_evidence"]
    with evidence_path.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    source_urls = {
        "TMview": "https://www.tmdn.org/tmview/",
        "UKIPO": "https://trademarks.ipo.gov.uk/ipo-tmtext",
        "USPTO": "https://tmsearch.uspto.gov/",
    }
    for source, source_url in source_urls.items():
        for suffix, query in (("exact", "conflict"), ("near", "conflict [fuzzy-1]")):
            rows.append(
                {
                    "evidence_id": f"history-{source}-{suffix}",
                    "domain": "conflict.si",
                    "source": source,
                    "jurisdiction": "synthetic history",
                    "query": query,
                    "utc_timestamp": "2026-10-13T05:00:00Z",
                    "result_status": "blocked" if source == "UKIPO" else "unverified",
                    "source_url": source_url,
                    "result_url": "",
                    "mark_number": "",
                    "mark_name": "",
                    "mark_status": "",
                    "classes": "",
                    "notes": "Historical per-name coverage; no clearance implied.",
                    "proof_kind": "",
                    "proof_path": "",
                }
            )
    with evidence_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=TRADEMARK_EVIDENCE_FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    assert write_trademark_summary(
        evidence_path,
        paths["shortlist"],
        paths["trademark_summary"],
        top_n=100,
        candidates_path=paths["candidates"],
    ) == 2


def test_delivery_validation_passes_honest_blocked_workbook_and_says_so(tmp_path, capsys):
    paths = make_fixture(tmp_path)
    _add_prior_screen_record(paths)
    workbook = tmp_path / "ranked.xlsx"
    export_workbook(paths, workbook)

    report = validate_delivery(workbook, paths)
    assert report.candidate_count == 4
    assert report.shortlist_count == 1
    assert report.research_unresolved is True
    assert report.sheet_rows["Unverified"] == 0

    assert main(
        ["validate-delivery", "--workbook", str(workbook), *_cli_inputs(paths)]
    ) == 0
    output = capsys.readouterr().out
    assert "Structural delivery PASS" in output
    assert "Research screening BLOCKED" in output
    assert "Project completion not certified" in output


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        ("rank", "Shortlist!A2"),
        ("score", "Scores!"),
        ("source_row", "All candidates!"),
        ("nonshortlist_registered", "All candidates!"),
        ("status", "Shortlist!"),
        ("evidence_row", "RDAP evidence!"),
        ("methodology_hash", "Methodology!"),
        ("missing_sheet", "sheet list differs"),
        ("formula", "formula cell is forbidden"),
        ("money_blank", "Shortlist!"),
    ],
)
def test_delivery_validator_rejects_workbook_tampering(tmp_path, mutation, message):
    paths = make_fixture(tmp_path)
    original = tmp_path / "ranked.xlsx"
    export_workbook(paths, original)
    tampered = tmp_path / f"{mutation}.xlsx"
    shutil.copyfile(original, tampered)
    workbook = load_workbook(tampered)

    if mutation == "rank":
        workbook["Shortlist"]["A2"] = 2
    elif mutation == "score":
        headers = [cell.value for cell in workbook["Scores"][1]]
        workbook["Scores"].cell(2, headers.index("total") + 1).value = 1.0
    elif mutation == "source_row":
        headers = [cell.value for cell in workbook["All candidates"][1]]
        workbook["All candidates"].cell(2, headers.index("tags") + 1).value = "changed"
    elif mutation == "nonshortlist_registered":
        sheet = workbook["All candidates"]
        headers = [cell.value for cell in sheet[1]]
        domain_column = headers.index("domain") + 1
        status_column = headers.index("latest_rdap_classification") + 1
        row = next(
            index for index in range(2, sheet.max_row + 1)
            if sheet.cell(index, domain_column).value == "taken.si"
        )
        sheet.cell(row, status_column).value = "potentially_unregistered"
    elif mutation == "status":
        headers = [cell.value for cell in workbook["Shortlist"][1]]
        workbook["Shortlist"].cell(2, headers.index("trademark_status") + 1).value = "cleared"
    elif mutation == "evidence_row":
        headers = [cell.value for cell in workbook["RDAP evidence"][1]]
        workbook["RDAP evidence"].cell(2, headers.index("classification") + 1).value = "registered"
    elif mutation == "methodology_hash":
        sheet = workbook["Methodology"]
        headers = [cell.value for cell in sheet[1]]
        topic_column = headers.index("Topic") + 1
        hash_column = headers.index("SHA-256") + 1
        row = next(
            index for index in range(2, sheet.max_row + 1)
            if sheet.cell(index, topic_column).value == "Source CSV: candidates"
        )
        sheet.cell(row, hash_column).value = "0" * 64
    elif mutation == "missing_sheet":
        del workbook["Scores"]
    elif mutation == "formula":
        workbook["Shortlist"]["C2"] = "=1+1"
    elif mutation == "money_blank":
        headers = [cell.value for cell in workbook["Shortlist"][1]]
        workbook["Shortlist"].cell(2, headers.index("valuation_low") + 1).value = 0
    workbook.save(tampered)
    workbook.close()

    with pytest.raises(ValueError, match=message):
        validate_delivery(tampered, paths)
