import json
import re
from html.parser import HTMLParser
from pathlib import Path

import openpyxl

REPORT = Path(__file__).resolve().parents[1] / "output" / "si_domains_ranked.html"


class ReportParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_summary = False
        self.in_owned_name = False
        self.in_owned_list_item = False
        self.in_data_script = False
        self.in_script = False
        self.in_owned_chip = False
        self.owned_names = []
        self.owned_items = []
        self.current_item = []
        self.data_json = ""
        self.script_text = []
        self.owned_chip_count = ""
        self.has_owned_chip = False
        self.in_owned_chip_count = False

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        classes = values.get("class", "").split()
        if tag == "section" and values.get("id") == "owned-summary":
            self.in_summary = True
        if tag == "li" and self.in_summary:
            self.in_owned_list_item = True
            self.current_item = []
        if tag == "span" and self.in_summary and "owned-name" in classes:
            self.in_owned_name = True
        if tag == "script":
            self.in_script = True
            self.in_data_script = values.get("id") == "data"
        if tag == "button" and values.get("data-owned") == "true":
            self.has_owned_chip = True
            self.in_owned_chip = True
        if tag == "i" and self.in_owned_chip:
            self.in_owned_chip_count = True

    def handle_data(self, data):
        if self.in_owned_name:
            self.owned_names.append(data.strip())
        if self.in_owned_list_item:
            self.current_item.append(data.strip())
        if self.in_data_script:
            self.data_json += data
        if self.in_script:
            self.script_text.append(data)
        if self.in_owned_chip_count:
            self.owned_chip_count += data.strip()

    def handle_endtag(self, tag):
        if tag == "span" and self.in_owned_name:
            self.in_owned_name = False
        if tag == "li" and self.in_owned_list_item:
            self.owned_items.append(" ".join(filter(None, self.current_item)))
            self.in_owned_list_item = False
        if tag == "section" and self.in_summary:
            self.in_summary = False
        if tag == "script":
            self.in_script = False
            self.in_data_script = False
        if tag == "button" and self.in_owned_chip:
            self.in_owned_chip = False
        if tag == "i" and self.in_owned_chip_count:
            self.in_owned_chip_count = False


def test_report_contains_owned_chip_summary_badges_and_purchase_wording():
    parser = ReportParser()
    source = REPORT.read_text(encoding="ascii")
    parser.feed(source)

    assert parser.has_owned_chip
    assert parser.owned_chip_count == "10"
    assert len(parser.owned_items) == 10
    assert len(parser.owned_names) == 10
    assert all("Rank" in item and "Score" in item for item in parser.owned_items)
    assert "owned-badge" in "".join(parser.script_text)

    data = json.loads(parser.data_json)
    owned = [row for row in data["s"] if row[4]]
    assert len(owned) == 10
    assert {row[1] for row in owned} == set(parser.owned_names)
    assert all(row[5] == "owned by you (bought 2026-10-14)" for row in owned)

    assert 'role="tablist"' in source
    assert all(f'data-mode="{mode}">{label}</button>' in source for mode, label in (
        ("top", "Top 500"), ("available", "All available"), ("unavailable", "Unavailable")
    ))
    stats = re.findall(r'<div class="stat"><b>(.*?)</b><span>(.*?)</span>', source)
    assert stats == [
        ('6,204', 'checked'), ('3,385', 'available'), ('10', 'owned (yours)'),
        ('500', 'top 500 ranked'), ('2,809', 'unavailable'),
    ]
    owned_note = "Your 10 owned domains are not listed here; see Your domains and Top 500."
    available_panel = source.split('id="panel-available"', 1)[1].split("</section>", 1)[0]
    assert owned_note in available_panel
    assert source.count(owned_note) == 1
    assert len(data["a"]) == 3385
    assert len(data["u"]) == 2809
    owned_domains = {row[1] for row in data["s"] if row[4]}
    assert not any(row[0] in owned_domains for row in data["a"])
    assert sum(row[4] is not None for row in data["a"]) == 490
    assert all(not row[6] and row[8] == "404" and row[9] == "potentially_unregistered"
               and row[11] == "not_registered" for row in data["a"])
    assert all(row[4] is not None or row[5] for row in data["a"])
    assert all(row[3] and row[6] for row in data["u"])
    assert "const PAGE_SIZE = 50" in "".join(parser.script_text)


def test_report_has_three_tabpanels_with_counts_matching_workbook_data():
    parser = ReportParser()
    source = REPORT.read_text(encoding="ascii")
    parser.feed(source)
    data = json.loads(parser.data_json)

    panels = re.findall(r'<(?:div|section) id="panel-(top|available|unavailable)" role="tabpanel"', source)
    assert panels == ["top", "available", "unavailable"]
    tabs = re.findall(r'<button[^>]+role="tab"[^>]+aria-controls="panel-[^"]+"', source)
    assert len(tabs) == 3

    stats = {label: int(value.replace(",", ""))
             for value, label in re.findall(r'<div class="stat"><b>(.*?)</b><span>(.*?)</span>', source)}
    workbook = openpyxl.load_workbook(REPORT.with_name("si_domains_ranked.xlsx"), read_only=True, data_only=True)
    try:
        checked = workbook["All candidates"].max_row - 1
        shortlist = workbook["Shortlist"].max_row - 1
    finally:
        workbook.close()

    assert stats == {"checked": checked, "available": len(data["a"]), "owned (yours)": 10,
                     "top 500 ranked": shortlist, "unavailable": len(data["u"])}
    assert stats["checked"] == stats["available"] + stats["owned (yours)"] + stats["unavailable"]
    assert stats["top 500 ranked"] == len(data["s"]) == 500
    assert sum(row[4] is not None for row in data["a"]) == 490


def test_unavailable_evidence_and_available_exclusion_reasons_are_recorded():
    parser = ReportParser()
    parser.feed(REPORT.read_text(encoding="ascii"))
    data = json.loads(parser.data_json)
    workbook = openpyxl.load_workbook(REPORT.with_name("si_domains_ranked.xlsx"), read_only=True, data_only=True)
    try:
        sheet = workbook["All candidates"]
        headers = next(sheet.iter_rows(values_only=True))
        source_rows = {str(row[0]): dict(zip(headers, row))
                       for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None}
    finally:
        workbook.close()

    assert all(row[3] and row[6] for row in data["u"])
    for row in data["u"]:
        source = source_rows[row[0]]
        assert row[6] in {str(source.get("latest_rdap_timestamp") or ""),
                          str(source.get("latest_whois_timestamp") or "")}

    for row in data["a"]:
        if row[4] is not None:
            continue
        recorded = str(source_rows[row[0]].get("exclusion_reason") or "").strip()
        assert row[5] == (recorded or "not recorded")
