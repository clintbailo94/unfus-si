import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from si_domains.candidates.schema import CSV_FIELDS as CANDIDATE_FIELDS
from si_domains.research import write_opportunities
from si_domains.scoring import SCORE_FIELDS, score_domain
from si_domains.shortlist import build_outputs
from si_domains.trademarks import TRADEMARK_EVIDENCE_FIELDS, write_trademark_summary
from si_domains.valuation import (
    VALUATION_EVIDENCE_FIELDS,
    write_valuations,
)
from si_domains.delivery import validate_delivery
from si_domains.owned import OwnedDomain, load_owned_domains
from si_domains.workbook import _validate_registry_rows, build_workbook_model, export_workbook


RDAP_FIELDS = [
    "domain",
    "category",
    "http_status",
    "utc_timestamp",
    "attempts",
    "classification",
    "error",
]
WHOIS_FIELDS = [
    "domain",
    "category",
    "utc_timestamp",
    "classification",
    "matched_pattern",
    "error",
]


def _write_csv(path: Path, fields, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _read_csv(path: Path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def test_owned_domain_input_parses_and_rejects_invalid_rows(tmp_path):
    path = tmp_path / "owned_domains.txt"
    path.write_text(
        "# domain,purchase_date\nBright.SI,2026-10-14\n",
        encoding="ascii",
        newline="\n",
    )
    owners = load_owned_domains(
        path,
        candidate_domains={"bright.si"},
        shortlist_domains={"bright.si"},
    )
    assert owners == (OwnedDomain("bright.si", "2026-10-14"),)
    assert owners[0].availability_text == "owned by you (bought 2026-10-14)"

    invalid_rows = (
        ("bright.si,2026-10-14\nbright.si,2026-10-14\n", "duplicate owned domain"),
        ("bright.si,14-10-2026\n", "invalid purchase date"),
        ("unknown.si,2026-10-14\n", "not in data/candidates.csv"),
        ("bright.si,2026-10-14\n", "not in the 500-row shortlist"),
    )
    for content, message in invalid_rows:
        path.write_text(content, encoding="ascii", newline="\n")
        with pytest.raises(ValueError, match=message):
            load_owned_domains(
                path,
                candidate_domains={"bright.si"},
                shortlist_domains=set() if "shortlist" in message else {"bright.si"},
            )


def make_fixture(tmp_path: Path, *, hostile_text=" \\t=1+1", second_eligible=False, resume_evidence=False):
    """Build a complete, tiny offline package from synthetic rows."""
    paths = {name: tmp_path / f"{name}.csv" for name in (
        "candidates",
        "reserved_names",
        "restricted_names",
        "rdap_evidence",
        "whois_evidence",
        "scores",
        "shortlist",
        "excluded",
        "opportunities",
        "valuations",
        "valuation_evidence",
        "trademark_summary",
        "trademark_evidence",
        "sample",
    )}
    paths["owned_domains"] = tmp_path / "owned_domains.txt"
    paths["owned_domains"].write_text("# domain,purchase_date\n", encoding="ascii", newline="\n")
    candidates = [
        {"domain_label": "bright", "category": "technology", "method": "concept", "seed": "bright", "tags": hostile_text, "quality_score": "80"},
        {"domain_label": "taken", "category": "technology", "method": "concept", "seed": "taken", "tags": "", "quality_score": "80"},
        {"domain_label": "conflict", "category": "technology", "method": "concept", "seed": "conflict", "tags": "", "quality_score": "80"},
        {"domain_label": "si", "category": "technology", "method": "concept", "seed": "si", "tags": "", "quality_score": "80"},
    ]
    if second_eligible:
        candidates.append({"domain_label": "glow", "category": "technology", "method": "concept", "seed": "glow", "tags": "", "quality_score": "80"})
    _write_csv(paths["candidates"], list(CANDIDATE_FIELDS), candidates)
    _write_csv(
        paths["reserved_names"],
        ["label", "class", "source_url", "note"],
        [{"label": "si", "class": "reserved", "source_url": "https://www.register.si/en/reserved-domain-names/", "note": "Reserved for the Republic of Slovenia."}],
    )
    _write_csv(paths["restricted_names"], ["label", "class", "source_url", "note"], [])

    rdap_rows = [
        {"domain": "bright.si", "category": "technology", "http_status": "429", "utc_timestamp": "2026-10-13T01:00:00Z", "attempts": "1", "classification": "unverified", "error": "http_429"},
        {"domain": "bright.si", "category": "technology", "http_status": "404", "utc_timestamp": "2026-10-13T02:00:00Z", "attempts": "1", "classification": "potentially_unregistered", "error": ""},
        {"domain": "taken.si", "category": "technology", "http_status": "200", "utc_timestamp": "2026-10-13T02:01:00Z", "attempts": "1", "classification": "registered", "error": ""},
        {"domain": "conflict.si", "category": "technology", "http_status": "404", "utc_timestamp": "2026-10-13T02:02:00Z", "attempts": "1", "classification": "potentially_unregistered", "error": ""},
    ]
    if second_eligible:
        rdap_rows.append({"domain": "glow.si", "category": "technology", "http_status": "404", "utc_timestamp": "2026-10-13T02:03:00Z", "attempts": "1", "classification": "potentially_unregistered", "error": ""})
    _write_csv(paths["rdap_evidence"], RDAP_FIELDS, rdap_rows)
    whois_rows = [
        {"domain": "bright.si", "category": "technology", "utc_timestamp": "2026-10-13T03:00:00Z", "classification": "not_registered", "matched_pattern": "no_entries_found", "error": ""},
        {"domain": "conflict.si", "category": "technology", "utc_timestamp": "2026-10-13T03:01:00Z", "classification": "registered", "matched_pattern": "field:registrar", "error": ""},
    ]
    if second_eligible:
        whois_rows.append({"domain": "glow.si", "category": "technology", "utc_timestamp": "2026-10-13T03:02:00Z", "classification": "not_registered", "matched_pattern": "no_entries_found", "error": ""})
    _write_csv(paths["whois_evidence"], WHOIS_FIELDS, whois_rows)

    score_rows = []
    score_candidates = [candidates[0], candidates[2]]
    if second_eligible:
        score_candidates.append(candidates[-1])
    for candidate in score_candidates:
        score = score_domain(candidate)
        score_rows.append({field: str(score[field]) for field in SCORE_FIELDS})
    _write_csv(paths["scores"], list(SCORE_FIELDS), score_rows)

    excluded_path = paths["excluded"]
    build_outputs(
        paths["candidates"],
        paths["rdap_evidence"],
        paths["whois_evidence"],
        paths["scores"],
        paths["shortlist"],
        excluded_path,
        paths["sample"],
    )

    urls = {
        "TMview": "https://www.tmdn.org/tmview/",
        "UKIPO": "https://trademarks.ipo.gov.uk/ipo-tmtext",
        "USPTO": "https://tmsearch.uspto.gov/",
    }
    trademark_rows = []
    for source, url in urls.items():
        trademark_rows.append({
            "evidence_id": f"fixture-{source}",
            "domain": "*",
            "source": source,
            "jurisdiction": "synthetic fixture",
            "query": "source access",
            "utc_timestamp": "2026-10-13T04:00:00Z",
            "result_status": "blocked" if source == "UKIPO" else "unverified",
            "source_url": url,
            "result_url": "",
            "mark_number": "",
            "mark_name": "",
            "mark_status": "",
            "classes": "",
            "notes": "No per-domain query ran; synthetic fixture.",
            "proof_kind": "",
            "proof_path": "",
        })
    if resume_evidence:
        trademark_rows.extend([
            {
                "evidence_id": "RESUME-USPTO-CONTROL",
                "domain": "*",
                "source": "USPTO",
                "jurisdiction": "US",
                "query": "google",
                "utc_timestamp": "2026-10-14T00:10:00Z",
                "result_status": "possible_conflict",
                "source_url": "https://tmsearch.uspto.gov/prod-stage-v1-0-0/tmsearch",
                "result_url": "https://tmsearch.uspto.gov/search/search-results?query=google",
                "mark_number": "12345678",
                "mark_name": "GOOGLE",
                "mark_status": "registered and renewed; live",
                "classes": "35",
                "notes": "Control query; HTTP 200 JSON; mode=exact; term=google; totalValue=2; first hit facts only",
                "proof_kind": "api_query_summary",
                "proof_path": "",
            },
            {
                "evidence_id": "RESUME-USPTO-00001",
                "domain": "probe.si",
                "source": "USPTO",
                "jurisdiction": "US",
                "query": "probe exact",
                "utc_timestamp": "2026-10-14T00:10:05Z",
                "result_status": "unverified",
                "source_url": "https://tmsearch.uspto.gov/prod-stage-v1-0-0/tmsearch",
                "result_url": "",
                "mark_number": "",
                "mark_name": "",
                "mark_status": "",
                "classes": "",
                "notes": "HTTP 200 JSON response was received; local evidence write failed",
                "proof_kind": "api_query_summary",
                "proof_path": "",
            },
        ])
    _write_csv(paths["trademark_evidence"], list(TRADEMARK_EVIDENCE_FIELDS), trademark_rows)
    write_trademark_summary(
        paths["trademark_evidence"],
        paths["shortlist"],
        paths["trademark_summary"],
        candidates_path=paths["candidates"],
    )
    # The incomplete synthetic screen uses the unverified mapping score.
    _write_csv(paths["valuation_evidence"], list(VALUATION_EVIDENCE_FIELDS), [])
    write_valuations(paths["shortlist"], paths["valuation_evidence"], paths["valuations"])
    write_opportunities(
        paths["shortlist"], paths["scores"], paths["valuations"],
        paths["trademark_summary"], paths["opportunities"],
    )
    return paths


def test_workbook_exports_latest_joins_numeric_cells_and_header_only_unverified(tmp_path):
    paths = make_fixture(tmp_path, resume_evidence=True)
    output = tmp_path / "ranked.xlsx"

    model = build_workbook_model(paths)
    export_workbook(model, output)

    assert output.is_file()
    normal = load_workbook(output, data_only=False)
    read_only = load_workbook(output, read_only=True, data_only=False)
    data_only = load_workbook(output, data_only=True)
    required = {
        "Shortlist", "Owned", "All candidates", "Unverified", "Special classes", "Methodology",
        "Scores", "RDAP evidence", "WHOIS evidence", "Valuation evidence", "Trademark evidence",
    }
    assert set(normal.sheetnames) == required
    assert normal["Shortlist"].max_row == 2
    assert normal["All candidates"].max_row == 5
    assert normal["RDAP evidence"].max_row == 5  # history is retained
    assert normal["Unverified"].max_row == 1

    methodology = normal["Methodology"]
    methodology_headers = [cell.value for cell in methodology[1]]
    methodology_columns = {header: index + 1 for index, header in enumerate(methodology_headers)}
    availability = next(
        row for row in methodology.iter_rows(min_row=2, values_only=True)
        if row[methodology_columns["Topic"] - 1] == "Availability: technology"
    )
    assert availability[methodology_columns["Details"] - 1] == "technology | 4 | 1 | 1 | 1 | 1"
    assert any(
        row[methodology_columns["Topic"] - 1] == "Demand root: pay"
        for row in methodology.iter_rows(min_row=2, values_only=True)
    )
    methodology_rows = list(methodology.iter_rows(min_row=2, values_only=True))
    manual_checks = next(
        row for row in methodology_rows
        if row[methodology_columns["Topic"] - 1] == "Manual trademark checks the owner can do"
    )
    assert "https://www.tmdn.org/tmview/" in manual_checks[methodology_columns["Details"] - 1]
    assert "REWORK1-USPTO-00259" in manual_checks[methodology_columns["Details"] - 1]
    assert "TMview source-wide refused: connection closed (evidence REWORK1-TMVIEW-PROBES, 2026-10-13T18:02:42Z); per-name results are unverified because of that refusal." in manual_checks[methodology_columns["Details"] - 1]
    assert "UKIPO source-wide refused: automated access refused (evidence PKGD-0003, 2026-10-13T17:09:52Z); per-name results are blocked because of that refusal." in manual_checks[methodology_columns["Details"] - 1]
    assert "PKG-G originally made 117 USPTO POSTs (one control and 116 candidates); T-18 added six (one control and five candidates)." in manual_checks[methodology_columns["Details"] - 1]
    completion_mode = next(
        row for row in methodology_rows
        if row[methodology_columns["Topic"] - 1] == "Completion-mode research validation"
    )
    assert "--completion-mode" in completion_mode[methodology_columns["Details"] - 1]
    assert "Missing pairs fail" in completion_mode[methodology_columns["Details"] - 1]
    assert "local evidence-write failure" in completion_mode[methodology_columns["Details"] - 1]
    assert "documented three-round limit" in completion_mode[methodology_columns["Details"] - 1]
    probe_outcome = next(
        row for row in methodology_rows
        if row[methodology_columns["Topic"] - 1] == "Trademark probe outcome"
    )
    probe_details = probe_outcome[methodology_columns["Details"] - 1]
    assert "2 USPTO POST attempts: 1 google control and 1 candidate queries" in probe_details
    assert "2 returned HTTP 200 JSON, 0 returned HTTP 403, and 0 retries" in probe_details
    assert "1 historical response row records a local evidence-write failure" in probe_details

    shortlist = normal["Shortlist"]
    headers = [cell.value for cell in shortlist[1]]
    column = {header: index + 1 for index, header in enumerate(headers)}
    assert isinstance(shortlist.cell(2, column["rank"]).value, int)
    assert isinstance(shortlist.cell(2, column["total"]).value, (int, float))
    assert isinstance(shortlist.cell(2, column["brandability"]).value, int)
    assert shortlist.cell(2, column["valuation_low"]).value is None
    assert shortlist.cell(2, column["valuation_high"]).value is None
    assert shortlist.cell(2, column["rdap_classification"]).value == "potentially_unregistered"
    assert shortlist.cell(2, column["whois_classification"]).value == "not_registered"

    all_headers = [cell.value for cell in normal["All candidates"][1]]
    all_column = {header: index + 1 for index, header in enumerate(all_headers)}
    bright = next(
        row for row in normal["All candidates"].iter_rows(min_row=2, values_only=True)
        if row[all_column["domain"] - 1] == "bright.si"
    )
    assert bright[all_column["latest_rdap_classification"] - 1] == "potentially_unregistered"
    assert bright[all_column["latest_rdap_http_status"] - 1] == 404
    conflict = next(
        row for row in normal["All candidates"].iter_rows(min_row=2, values_only=True)
        if row[all_column["domain"] - 1] == "conflict.si"
    )
    assert conflict[all_column["registry_status"] - 1] == "whois_registered_conflict"
    assert conflict[all_column["exclusion_reason"] - 1] == "whois_registered_conflict"

    tags_cell = normal["All candidates"].cell(2, all_column["tags"])
    assert tags_cell.value == " \\t=1+1"
    assert tags_cell.data_type == "s"
    assert "si.si" in {
        normal["Special classes"].cell(row, 2).value
        for row in range(2, normal["Special classes"].max_row + 1)
    }
    assert normal["Special classes"].freeze_panes
    assert not any(cell.data_type == "f" for sheet in normal for row in sheet.iter_rows() for cell in row)
    assert not normal._external_links
    assert read_only["Shortlist"].cell(2, column["domain"]).value == "bright.si"
    assert data_only["Shortlist"].cell(2, column["domain"]).value == "bright.si"

    normal.close()
    read_only.close()
    data_only.close()


def test_workbook_relabels_owned_registry_status_and_adds_owned_sheet(tmp_path):
    paths = make_fixture(tmp_path)
    paths["owned_domains"].write_text(
        "bright.si,2026-10-14\n", encoding="ascii", newline="\n"
    )
    output = tmp_path / "owned.xlsx"
    export_workbook(paths, output)

    workbook = load_workbook(output, data_only=True)
    try:
        shortlist = workbook["Shortlist"]
        shortlist_headers = {cell.value: index for index, cell in enumerate(shortlist[1])}
        row = next(row for row in shortlist.iter_rows(min_row=2, values_only=True) if row[shortlist_headers["domain"]] == "bright.si")
        assert row[shortlist_headers["owned"]] == "yes"
        assert row[shortlist_headers["rdap_classification"]] == "potentially_unregistered"
        assert row[shortlist_headers["whois_classification"]] == "not_registered"

        all_candidates = workbook["All candidates"]
        all_headers = {cell.value: index for index, cell in enumerate(all_candidates[1])}
        row = next(row for row in all_candidates.iter_rows(min_row=2, values_only=True) if row[all_headers["domain"]] == "bright.si")
        assert row[all_headers["owned"]] == "yes"
        assert row[all_headers["registry_status"]] == "owned by you (bought 2026-10-14)"
        assert row[all_headers["latest_rdap_classification"]] == "potentially_unregistered"
        assert row[all_headers["latest_whois_classification"]] == "not_registered"

        for sheet_name, source_name in (("RDAP evidence", "rdap_evidence"), ("WHOIS evidence", "whois_evidence")):
            evidence_sheet = workbook[sheet_name]
            evidence_headers = {cell.value: index for index, cell in enumerate(evidence_sheet[1])}
            evidence_rows = [
                row[evidence_headers["classification"]]
                for row in evidence_sheet.iter_rows(min_row=2, values_only=True)
                if row[evidence_headers["domain"]] == "bright.si"
            ]
            recorded_rows = [row["classification"] for row in _read_csv(paths[source_name]) if row["domain"] == "bright.si"]
            assert evidence_rows == recorded_rows

        owned_sheet = workbook["Owned"]
        assert owned_sheet.max_row == 2
        owned_headers = {cell.value: index for index, cell in enumerate(owned_sheet[1])}
        owned_row = next(owned_sheet.iter_rows(min_row=2, values_only=True))
        assert owned_row[owned_headers["domain"]] == "bright.si"
        assert owned_row[owned_headers["purchase_date"]] == "2026-10-14"
        assert isinstance(owned_row[owned_headers["rank"]], int)
        assert isinstance(owned_row[owned_headers["total"]], (int, float))
    finally:
        workbook.close()


def test_production_owned_sheet_has_ten_shortlist_rows_and_matching_ranks(tmp_path):
    output = tmp_path / "ranked.xlsx"
    export_workbook(build_workbook_model(), output)
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        shortlist = list(workbook["Shortlist"].iter_rows(values_only=True))
        shortlist_headers = {name: index for index, name in enumerate(shortlist[0])}
        shortlist_by_domain = {
            row[shortlist_headers["domain"]]: row for row in shortlist[1:]
        }
        all_candidates = list(workbook["All candidates"].iter_rows(values_only=True))
        all_headers = {name: index for index, name in enumerate(all_candidates[0])}
        all_by_domain = {
            row[all_headers["domain"]]: row for row in all_candidates[1:]
        }
        owned = list(workbook["Owned"].iter_rows(values_only=True))
        owned_headers = {name: index for index, name in enumerate(owned[0])}
        assert len(owned) - 1 == 10

        def latest_classifications(path):
            latest = {}
            for evidence in _read_csv(path):
                domain = evidence["domain"]
                if domain not in latest or evidence["utc_timestamp"] > latest[domain]["utc_timestamp"]:
                    latest[domain] = evidence
            return {domain: row["classification"] for domain, row in latest.items()}

        root = Path(__file__).resolve().parents[1]
        rdap = latest_classifications(root / "output" / "rdap_evidence.csv")
        whois = latest_classifications(root / "output" / "whois_evidence.csv")
        for row in owned[1:]:
            domain = row[owned_headers["domain"]]
            shortlist_row = shortlist_by_domain[domain]
            candidate_row = all_by_domain[domain]
            assert row[owned_headers["rank"]] == shortlist_row[shortlist_headers["rank"]]
            assert row[owned_headers["total"]] == shortlist_row[shortlist_headers["total"]]
            assert row[owned_headers["category"]] == shortlist_row[shortlist_headers["category"]]
            assert shortlist_row[shortlist_headers["owned"]] == "yes"
            assert shortlist_row[shortlist_headers["rdap_classification"]] == rdap[domain]
            assert shortlist_row[shortlist_headers["whois_classification"]] == whois[domain]
            assert candidate_row[all_headers["owned"]] == "yes"
            assert candidate_row[all_headers["latest_rdap_classification"]] == rdap[domain]
            assert candidate_row[all_headers["latest_whois_classification"]] == whois[domain]
            assert candidate_row[all_headers["registry_status"]] == f"owned by you (bought {row[owned_headers['purchase_date']]})"
    finally:
        workbook.close()


@pytest.mark.parametrize(
    "hostile",
    ["=1+1", "+1+2", "-1+2", "@SUM(A1:A2)", " \t=1+1", "\t@SUM(A1:A2)"],
)
def test_external_csv_text_with_formula_prefixes_stays_plain_text(tmp_path, hostile):
    paths = make_fixture(tmp_path, hostile_text=hostile)
    output = tmp_path / "ranked.xlsx"
    export_workbook(build_workbook_model(paths), output)

    workbook = load_workbook(output, data_only=False)
    headers = [cell.value for cell in workbook["All candidates"][1]]
    tags_column = headers.index("tags") + 1
    cell = next(
        workbook["All candidates"].cell(row, tags_column)
        for row in range(2, workbook["All candidates"].max_row + 1)
        if workbook["All candidates"].cell(row, headers.index("domain") + 1).value == "bright.si"
    )
    assert cell.value == hostile
    assert cell.data_type == "s"
    assert cell.hyperlink is None
    workbook.close()


def test_export_never_overwrites_a_source_csv(tmp_path):
    paths = make_fixture(tmp_path)
    original = paths["candidates"].read_bytes()
    model = build_workbook_model(paths)

    with pytest.raises(ValueError, match="cannot overwrite a source CSV"):
        export_workbook(model, paths["candidates"])

    assert paths["candidates"].read_bytes() == original


def test_shuffled_opportunities_export_in_rank_order_and_validator_rejects_shuffled_rows(tmp_path):
    paths = make_fixture(tmp_path, second_eligible=True)
    opportunity_rows = _read_csv(paths["opportunities"])
    assert len(opportunity_rows) == 2
    _write_csv(paths["opportunities"], opportunity_rows[0].keys(), reversed(opportunity_rows))
    output = tmp_path / "ranked.xlsx"

    export_workbook(paths, output)
    workbook = load_workbook(output)
    shortlist = workbook["Shortlist"]
    headers = [cell.value for cell in shortlist[1]]
    rank_column = headers.index("rank") + 1
    ranks = [shortlist.cell(row, rank_column).value for row in range(2, shortlist.max_row + 1)]
    assert ranks == [1, 2]
    workbook.close()

    tampered = tmp_path / "shuffled.xlsx"
    workbook = load_workbook(output)
    sheet = workbook["Shortlist"]
    first = [sheet.cell(2, column).value for column in range(1, sheet.max_column + 1)]
    second = [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)]
    for column, value in enumerate(second, start=1):
        sheet.cell(2, column).value = value
    for column, value in enumerate(first, start=1):
        sheet.cell(3, column).value = value
    workbook.save(tampered)
    workbook.close()
    with pytest.raises(ValueError, match="Shortlist!A2"):
        validate_delivery(tampered, paths)


@pytest.mark.parametrize(
    ("hostile_text", "message"),
    [("x" * 32768, "32767-character cell limit"), ("\x00", "invalid XML character")],
    ids=("over-limit", "xml-control"),
)
def test_unrepresentable_external_text_fails_before_replacing_existing_workbook(
    tmp_path, hostile_text, message
):
    paths = make_fixture(tmp_path, hostile_text=hostile_text)
    output = tmp_path / "ranked.xlsx"
    previous = b"previous delivery must survive invalid input"
    output.write_bytes(previous)
    source_bytes = paths["candidates"].read_bytes()

    with pytest.raises(ValueError, match=message):
        export_workbook(paths, output)

    assert output.read_bytes() == previous
    assert paths["candidates"].read_bytes() == source_bytes


def test_registry_evidence_keeps_historical_category_after_curated_name_retag():
    candidates = {
        "brena.si": {
            "domain_label": "brena",
            "category": "invented",
            "method": "given-name-reclassified",
            "tags": "phonotactic;given_name",
        }
    }
    rdap_rows = [
        {
            "domain": "brena.si",
            "category": "dictionary",
            "http_status": "404",
            "utc_timestamp": "2026-10-13T01:00:00Z",
            "attempts": "1",
            "classification": "potentially_unregistered",
            "error": "",
        }
    ]
    whois_rows = [
        {
            "domain": "brena.si",
            "category": "dictionary",
            "utc_timestamp": "2026-10-13T02:00:00Z",
            "classification": "not_registered",
            "matched_pattern": "no_entries_found",
            "error": "",
        }
    ]

    rdap, whois = _validate_registry_rows(candidates, rdap_rows, whois_rows)

    assert rdap["brena.si"]["category"] == "dictionary"
    assert whois["brena.si"]["category"] == "dictionary"
